# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "beautifulsoup4>=4.12",
#   "pillow>=10.0",
#   "numpy>=1.26",
#   "python-docx>=1.1",
#   "openpyxl>=3.1",
#   "python-pptx>=1.0",
#   "pypdf>=4.0",
# ]
# ///
"""Build a searchable source library from the Evolve "Origins Program" course folder.

Every stage is resumable: it records progress in _pipeline/state and skips work that is
already complete. Original course files are only ever read, never modified.

Stages
  inventory   parse lesson pages, match videos and templates          free, local
  prepare     extract audio, detect silence, plan + cut audio chunks   free, local
  transcribe  transcribe each audio chunk with Gemini                  PAID API
  assemble    merge chunks -> transcript.json / .srt / .md             free, local
  frames      sample keyframes, pick visually distinct candidates      free, local
  visual      describe candidate frames against the transcript         PAID API
  summarize   per-lesson summaries (kept separate from transcripts)    PAID API
  templates   extract searchable text from course templates            free, local
  index       INDEX.md, PROCESSING_REPORT.md, search.html               free, local
  follow      loop: assemble + visual + summarize lessons as they become ready

Usage
  uv run --script _pipeline/origins_pipeline.py <stage> [--lessons 20,25] [--workers 4]
  uv run --script _pipeline/origins_pipeline.py estimate
  uv run --script _pipeline/origins_pipeline.py status
  uv run --script _pipeline/origins_pipeline.py all --lessons 20
"""
from __future__ import annotations

import argparse
import base64
import concurrent.futures as cf
import datetime as dt
import glob
import hashlib
import html
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import sys
import threading
import time
import traceback
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
LIB = HERE.parent

DEFAULT_CONFIG = {
    "course_dir": str(LIB.parent / "Origins Program ($0-$100k mo)"),
    "cache_dir": str(Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "OriginsSourceLibrary" / "cache"),
    "course_name": "Origins Program ($0-$100k/mo)",
    "transcribe_model": "gemini-3.8-flash",
    "fallback_model": "gemini-2.5-flash",
    "visual_model": "gemini-3.8-flash",
    "summary_model": "gemini-3.8-flash",
    "thinking_level": "low",
    "chunk_target_s": 300,
    "chunk_window_s": 60,
    "silence_min_s": 0.35,
    "max_attempts": 3,
    "workers": 4,
    "visual_batch": 12,
    "max_cost_usd": 30.0,
    # USD per 1M tokens, standard paid tier, from ai.google.dev/gemini-api/docs/pricing (2026-10-06).
    "prices": {
        "gemini-3.8-flash": {"input": 0.75, "audio": 0.75, "output": 3.75},
        "gemini-2.5-flash": {"input": 0.30, "audio": 1.00, "output": 2.50},
    },
}


def load_config() -> dict:
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    p = HERE / "config.json"
    if p.exists():
        cfg.update(json.loads(p.read_text(encoding="utf-8")))
    return cfg


CFG = load_config()
COURSE = Path(CFG["course_dir"])
CACHE = Path(CFG["cache_dir"])
STATE = HERE / "state"
RAW = HERE / "raw"
LOGS = HERE / "logs"
OUT_LESSONS = LIB / "lessons"
OUT_TEMPLATES = LIB / "templates"
CATALOG = STATE / "catalog.json"
USAGE = STATE / "usage.jsonl"

LOCK = threading.RLock()


# --------------------------------------------------------------------------- utilities

def log(msg: str) -> None:
    line = f"[{dt.datetime.now():%H:%M:%S}] {msg}"
    with LOCK:
        try:
            print(line, flush=True)
        except UnicodeEncodeError:
            print(line.encode("ascii", "replace").decode(), flush=True)
        LOGS.mkdir(parents=True, exist_ok=True)
        with open(LOGS / "pipeline.log", "a", encoding="utf-8") as f:
            f.write(line + "\n")


def now_iso() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def _replace(tmp: Path, path: Path) -> None:
    for i in range(20):  # OneDrive can briefly lock files while syncing
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            time.sleep(0.25 * (i + 1))
    os.replace(tmp, path)


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8", newline="\n")
    _replace(tmp, path)


def write_json(path: Path, obj) -> None:
    write_text(path, json.dumps(obj, ensure_ascii=False, indent=2))


def read_json(path: Path, default=None):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def hms(t: float) -> str:
    t = max(0, int(round(t)))
    return f"{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}"


def srt_ts(t: float) -> str:
    ms = max(0, int(round(t * 1000)))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d},{ms % 1000:03d}"


def slugify(s: str, maxlen: int = 70) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return s[:maxlen].rstrip("-") or "untitled"


def rel_link(target: Path, start: Path) -> str:
    """Relative, URL-encoded markdown link from directory `start` to `target`."""
    r = os.path.relpath(target, start).replace("\\", "/")
    return urllib.parse.quote(r, safe="/.-_~")


def parse_ts(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = str(v).strip().replace(",", ".")
    try:
        nums = [float(p) for p in s.split(":")]
    except ValueError:
        return None
    if len(nums) == 1:
        return nums[0]
    if len(nums) == 2:
        return nums[0] * 60 + nums[1]
    if len(nums) == 3:
        return nums[0] * 3600 + nums[1] * 60 + nums[2]
    return None


def intervals_measure(iv: list[tuple[float, float]]) -> float:
    return sum(max(0.0, e - s) for s, e in iv)


def intervals_union(iv: list[tuple[float, float]]) -> list[tuple[float, float]]:
    out: list[list[float]] = []
    for s, e in sorted(iv):
        if out and s <= out[-1][1]:
            out[-1][1] = max(out[-1][1], e)
        else:
            out.append([s, e])
    return [(a, b) for a, b in out]


def intervals_subtract(a: list[tuple[float, float]], b: list[tuple[float, float]]) -> list[tuple[float, float]]:
    b = intervals_union(b)
    out = []
    for s, e in a:
        cur = s
        for bs, be in b:
            if be <= cur or bs >= e:
                continue
            if bs > cur:
                out.append((cur, bs))
            cur = max(cur, be)
        if cur < e:
            out.append((cur, e))
    return out


# --------------------------------------------------------------------------- ffmpeg

def find_tool(name: str) -> str:
    p = shutil.which(name)
    if p:
        return p
    pattern = os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\WinGet\Packages\Gyan.FFmpeg*\*\bin")
    for d in sorted(glob.glob(pattern)):
        cand = Path(d) / f"{name}.exe"
        if cand.exists():
            return str(cand)
    raise SystemExit(f"{name} not found. Install with: winget install Gyan.FFmpeg")


FFMPEG = find_tool("ffmpeg")
FFPROBE = find_tool("ffprobe")


def run(cmd: list[str], check: bool = True) -> subprocess.CompletedProcess:
    p = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if check and p.returncode != 0:
        raise RuntimeError(f"command failed ({p.returncode}): {' '.join(cmd[:6])} ...\n{p.stderr[-2000:]}")
    return p


def probe_duration(path: Path) -> float:
    p = run([FFPROBE, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)])
    return float(p.stdout.strip())


def probe_video(path: Path) -> dict:
    p = run([FFPROBE, "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height,codec_name",
             "-of", "json", str(path)])
    s = json.loads(p.stdout)["streams"][0]
    return {"width": s["width"], "height": s["height"], "codec": s["codec_name"]}


# --------------------------------------------------------------------------- state

def lesson_dir_state(lid: str) -> Path:
    return STATE / "lessons" / lid


def load_state(lid: str) -> dict:
    return read_json(lesson_dir_state(lid) / "state.json", {}) or {}


def update_state(lid: str, fn) -> dict:
    with LOCK:
        st = load_state(lid)
        fn(st)
        write_json(lesson_dir_state(lid) / "state.json", st)
        return st


def load_catalog() -> list[dict]:
    cat = read_json(CATALOG)
    if not cat:
        raise SystemExit("No catalog yet. Run the 'inventory' stage first.")
    return cat["lessons"]


def select_lessons(arg: str | None, video_only: bool = False) -> list[dict]:
    lessons = load_catalog()
    if arg:
        wanted = set()
        for part in arg.split(","):
            part = part.strip()
            if "-" in part:
                a, b = part.split("-")
                wanted.update(range(int(a), int(b) + 1))
            elif part:
                wanted.add(int(part))
        lessons = [L for L in lessons if L["num"] in wanted]
    if video_only:
        lessons = [L for L in lessons if L.get("video")]
    return lessons


# --------------------------------------------------------------------------- Gemini

class BudgetExceeded(Exception):
    pass


class QuotaExhausted(Exception):
    pass


def gemini_key() -> str:
    key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if not key and sys.platform == "win32":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as k:
                key = winreg.QueryValueEx(k, "GEMINI_API_KEY")[0]
        except OSError:
            key = None
    if not key:
        raise SystemExit("GEMINI_API_KEY is not set.")
    return key


def usage_total() -> float:
    if not USAGE.exists():
        return 0.0
    total = 0.0
    for line in USAGE.read_text(encoding="utf-8").splitlines():
        if line.strip():
            total += json.loads(line).get("cost_usd", 0.0)
    return total


def price_call(model: str, usage: dict) -> float:
    pr = CFG["prices"].get(model) or CFG["prices"]["gemini-3.8-flash"]
    audio = 0
    for d in usage.get("promptTokensDetails", []) or []:
        if d.get("modality") == "AUDIO":
            audio += d.get("tokenCount", 0)
    prompt = usage.get("promptTokenCount", 0)
    out = usage.get("candidatesTokenCount", 0) + usage.get("thoughtsTokenCount", 0)
    return ((prompt - audio) * pr["input"] + audio * pr["audio"] + out * pr["output"]) / 1e6


def gemini_generate(model: str, parts: list[dict], gen_config: dict, raw_path: Path, meta: dict) -> dict:
    """POST generateContent; the raw response (or error body) is written to raw_path before returning."""
    if usage_total() >= CFG["max_cost_usd"]:
        raise BudgetExceeded(f"Recorded spend reached max_cost_usd={CFG['max_cost_usd']}")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    body = {"contents": [{"role": "user", "parts": parts}], "generationConfig": gen_config}
    data = json.dumps(body).encode("utf-8")
    req_meta = {"model": model, "generationConfig": gen_config, "parts": [
        {k: (f"<{len(v.get('data', ''))} base64 chars, {v.get('mimeType')}>" if k == "inlineData" else v)
         for k, v in p.items()} for p in parts]}
    backoff = 10.0
    transient = 0
    while True:
        t0 = time.time()
        req = urllib.request.Request(url, data=data, method="POST", headers={
            "Content-Type": "application/json", "x-goog-api-key": gemini_key()})
        status, payload, err = None, None, None
        try:
            with urllib.request.urlopen(req, timeout=900) as r:
                status = r.status
                payload = json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            status = e.code
            txt = e.read().decode("utf-8", "replace")
            try:
                payload = json.loads(txt)
            except json.JSONDecodeError:
                payload = {"raw_error_text": txt}
            err = f"HTTP {e.code}"
        except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
            err = f"{type(e).__name__}: {e}"
        elapsed = round(time.time() - t0, 1)
        record = {"saved_at": now_iso(), "meta": meta, "request": req_meta, "http_status": status,
                  "elapsed_s": elapsed, "error": err, "response": payload}
        if status == 200:
            write_json(raw_path, record)
            usage = payload.get("usageMetadata", {})
            cost = price_call(model, usage)
            with LOCK:
                USAGE.parent.mkdir(parents=True, exist_ok=True)
                with open(USAGE, "a", encoding="utf-8") as f:
                    f.write(json.dumps({"ts": now_iso(), "model": model, **meta, "usage": usage,
                                        "cost_usd": round(cost, 6), "raw": str(raw_path.relative_to(HERE))}) + "\n")
            return payload
        # keep every failed response too, numbered so nothing is overwritten
        n = 1
        while (fail_path := raw_path.with_name(raw_path.stem + f".http-fail{n}.json")).exists():
            n += 1
        write_json(fail_path, record)
        msg = json.dumps(payload)[:600] if payload else err
        if status == 429:
            if "PerDay" in json.dumps(payload or {}):
                raise QuotaExhausted(f"Daily quota exhausted for {model}: {msg}")
            delay = backoff
            m = re.search(r'"retryDelay":\s*"(\d+(?:\.\d+)?)s"', json.dumps(payload or {}))
            if m:
                delay = float(m.group(1)) + 2
            log(f"  rate limited ({meta.get('ref')}); sleeping {delay:.0f}s")
            time.sleep(delay)
            backoff = min(backoff * 1.5, 120)
            continue
        if status is None or status >= 500:
            transient += 1
            if transient > 6:
                raise RuntimeError(f"Gemini transient failures exhausted: {msg}")
            log(f"  transient error ({meta.get('ref')}): {err}; retry in {backoff:.0f}s")
            time.sleep(backoff)
            backoff = min(backoff * 2, 240)
            continue
        if status == 400 and "thinking" in json.dumps(payload or {}).lower() and "thinkingConfig" in gen_config:
            log(f"  model rejected thinkingConfig; retrying without it")
            gen_config = {k: v for k, v in gen_config.items() if k != "thinkingConfig"}
            body["generationConfig"] = gen_config
            data = json.dumps(body).encode("utf-8")
            continue
        raise RuntimeError(f"Gemini request failed ({status}): {msg}")


def response_text(resp: dict) -> tuple[str | None, str]:
    cands = resp.get("candidates") or []
    if not cands:
        return None, (resp.get("promptFeedback") or {}).get("blockReason", "NO_CANDIDATES")
    c = cands[0]
    parts = (c.get("content") or {}).get("parts") or []
    text = "".join(p.get("text", "") for p in parts if not p.get("thought"))
    return text, c.get("finishReason", "UNKNOWN")


def loads_lenient(text: str):
    try:
        return json.loads(text)
    except (json.JSONDecodeError, TypeError):
        pass
    if not text:
        return None
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
    try:
        return json.loads(t)
    except json.JSONDecodeError:
        a, b = t.find("{"), t.rfind("}")
        if a >= 0 and b > a:
            try:
                return json.loads(t[a:b + 1])
            except json.JSONDecodeError:
                return None
    return None


# --------------------------------------------------------------------------- inventory

def html_to_md(node) -> str:
    from bs4 import NavigableString, Tag

    def rec(el) -> str:
        if isinstance(el, NavigableString):
            return re.sub(r"\s+", " ", str(el))
        if not isinstance(el, Tag):
            return ""
        name = el.name
        if name == "br":
            return "\n"
        inner = "".join(rec(c) for c in el.children)
        if name in ("strong", "b"):
            s = inner.strip()
            return f"**{s}**" if s else inner
        if name in ("em", "i"):
            s = inner.strip()
            return f"*{s}*" if s else inner
        if name == "a":
            href = el.get("href", "")
            txt = inner.strip() or href
            return f"[{txt}]({href})" if href else txt
        if name == "p":
            return inner.strip() + "\n\n"
        if re.fullmatch(r"h[1-6]", name):
            return "#" * (int(name[1]) + 1) + " " + inner.strip() + "\n\n"
        if name == "li":
            return "- " + inner.strip() + "\n"
        if name in ("ul", "ol"):
            return "\n" + inner + "\n"
        if name == "blockquote":
            return "> " + inner.strip() + "\n\n"
        if name == "img":
            return f"![{el.get('alt', '')}]({el.get('src', '')})"
        return inner

    md = rec(node)
    md = re.sub(r"[ \t]+\n", "\n", md)
    return re.sub(r"\n{3,}", "\n\n", md).strip()


def parse_nav(text: str) -> tuple[list[dict], list[dict]]:
    sections, modules = [], []
    for m in re.finditer(r'data-rbd-draggable-id="(setModule|set)-([0-9a-f]{32})(?:-([0-9a-f]{32}))?"', text):
        kind, a, b = m.groups()
        t = re.search(r'title="([^"]*)"', text[m.end():m.end() + 4000])
        title = html.unescape(t.group(1)) if t else ""
        if kind == "set":
            sections.append({"id": a, "title": title})
        else:
            modules.append({"md": a, "set": b, "title": title})
    return sections, modules


def norm_title(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower())


def stage_inventory(args) -> None:
    from bs4 import BeautifulSoup

    files = sorted(COURSE.iterdir())
    by_num: dict[int, dict[str, list[Path]]] = {}
    for f in files:
        m = re.match(r"^(\d+)\.\s", f.name)
        if not m or not f.is_file():
            continue
        by_num.setdefault(int(m.group(1)), {}).setdefault(f.suffix.lower(), []).append(f)

    nav_sections, nav_modules = [], []
    lessons = []
    for num in sorted(by_num):
        group = by_num[num]
        htmls = group.get(".html", [])
        page = htmls[0] if htmls else None
        title, desc_md, links, dur_label = None, "", [], None
        if page:
            text = page.read_text(encoding="utf-8", errors="replace")
            if not nav_modules:
                nav_sections, nav_modules = parse_nav(text)
            soup = BeautifulSoup(text, "html.parser")
            if soup.title:
                title = re.sub(r"\s+-\s+Origins Program.*$", "", soup.title.get_text()).strip()
            editor = soup.find("div", class_="ProseMirror")
            if editor:
                desc_md = html_to_md(editor)
                links = [{"text": a.get_text(" ", strip=True), "url": a["href"]}
                         for a in editor.find_all("a", href=True)]
            m = re.search(r'skool-mux-attachment.*?>(\d{1,2}:\d{2}(?::\d{2})?)</div>', text, re.S)
            dur_label = m.group(1) if m else None
        video = (group.get(".mp4") or [None])[0]
        if not title:
            base = (page or video).name
            title = re.sub(r"^\d+\.\s*", "", re.split(r"_Origins Program| Origins Program", base)[0])
        templates = [f.name for ext in (".docx", ".xlsx", ".pptx", ".pdf") for f in group.get(ext, [])]
        rec = {"num": num, "id": f"{num:03d}-{slugify(title)}", "title": title,
               "html": page.name if page else None, "video": video.name if video else None,
               "templates": templates, "links": links, "description_md": desc_md,
               "duration_label": dur_label}
        if video:
            rec["duration"] = round(probe_duration(video), 3)
            rec.update(probe_video(video))
        lessons.append(rec)

    # attach Skool section + module id by title match, falling back to course order
    set_title = {s["id"]: s["title"] for s in nav_sections}
    set_order = {s["id"]: i for i, s in enumerate(nav_sections)}
    by_title = {norm_title(m["title"]): m for m in nav_modules}
    for i, L in enumerate(lessons):
        mod = by_title.get(norm_title(L["title"])) or (nav_modules[L["num"] - 1] if L["num"] - 1 < len(nav_modules) else None)
        if mod:
            L["section"] = set_title.get(mod["set"], "")
            L["section_order"] = set_order.get(mod["set"], 99)
            L["skool_url"] = f"https://www.skool.com/evolve-8484/classroom/2580d099?md={mod['md']}"
        else:
            L["section"], L["section_order"] = "Unsorted", 99
    write_json(CATALOG, {"generated_at": now_iso(), "course_dir": str(COURSE),
                         "sections": [s["title"] for s in nav_sections], "lessons": lessons})
    nv = sum(1 for L in lessons if L["video"])
    log(f"inventory: {len(lessons)} lessons, {nv} with video, "
        f"{sum(len(L['templates']) for L in lessons)} templates, {len(nav_sections)} sections")


# --------------------------------------------------------------------------- prepare

def detect_silences(audio: Path, noise_db: float, min_s: float, duration: float) -> list[tuple[float, float]]:
    p = run([FFMPEG, "-hide_banner", "-nostats", "-i", str(audio), "-af",
             f"silencedetect=noise={noise_db:.1f}dB:d={min_s}", "-f", "null", "-"])
    out, start = [], None
    for line in p.stderr.splitlines():
        m = re.search(r"silence_start: (-?[\d.]+)", line)
        if m:
            start = max(0.0, float(m.group(1)))
        m = re.search(r"silence_end: ([\d.]+)", line)
        if m and start is not None:
            out.append((round(start, 3), round(float(m.group(1)), 3)))
            start = None
    if start is not None:
        out.append((round(start, 3), round(duration, 3)))
    return out


def mean_volume(audio: Path) -> float:
    p = run([FFMPEG, "-hide_banner", "-nostats", "-i", str(audio), "-af", "volumedetect", "-f", "null", "-"])
    m = re.search(r"mean_volume: (-?[\d.]+) dB", p.stderr)
    return float(m.group(1)) if m else -25.0


def quiet_point(audio: Path, lo: float, hi: float) -> tuple[float, float]:
    """Centre of the quietest 400 ms in [lo, hi] and how far (dB) it dips below the window's median level."""
    import numpy as np

    raw = run_bytes([FFMPEG, "-v", "error", "-ss", f"{lo:.3f}", "-t", f"{hi - lo:.3f}", "-i", str(audio),
                     "-f", "f32le", "-ac", "1", "-ar", "16000", "-"])
    x = np.frombuffer(raw, dtype=np.float32)
    hop = 800  # 50 ms
    n = len(x) // hop
    if n < 10:
        return (lo + hi) / 2, 0.0
    rms = np.sqrt((x[: n * hop].reshape(n, hop) ** 2).mean(axis=1) + 1e-12)
    smooth = np.convolve(rms, np.ones(8) / 8, mode="same")
    i = int(np.argmin(smooth[4:-4])) + 4
    db = 20 * np.log10(smooth + 1e-9)
    return lo + (i + 0.5) * hop / 16000, float(np.median(db) - db[i])


def run_bytes(cmd: list[str]) -> bytes:
    p = subprocess.run(cmd, capture_output=True)
    if p.returncode != 0:
        raise RuntimeError(p.stderr.decode("utf-8", "replace")[-1000:])
    return p.stdout


def plan_chunks(duration: float, silences: list[tuple[float, float]], target: float, window: float,
                audio: Path | None = None) -> list[dict]:
    chunks, start = [], 0.0
    while duration - start > target + window:
        cut, kind, dip = None, "hard", None
        for w in (window, window * 2):
            lo, hi = start + target - w, start + target + w
            best = None
            for s, e in silences:
                s2, e2 = max(s, lo), min(e, hi)
                if e2 - s2 < 0.3:
                    continue
                mid = (s2 + e2) / 2
                score = (e2 - s2) - 0.004 * abs(mid - (start + target))
                if best is None or score > best[0]:
                    best = (score, mid)
            if best:
                cut, kind = best[1], "silence"
                break
        if cut is None and audio is not None:
            # noisy/music bed: no true silence, so cut at the quietest moment (usually a breath between phrases)
            cut, dip = quiet_point(audio, start + target - window, start + target + window)
            kind = "quiet" if dip >= 6 else "hard"
        if cut is None:
            cut = start + target
        c = {"idx": len(chunks), "start": round(start, 3), "end": round(cut, 3), "end_boundary": kind}
        if dip is not None:
            c["dip_db"] = round(dip, 1)
        chunks.append(c)
        start = cut
    chunks.append({"idx": len(chunks), "start": round(start, 3), "end": round(duration, 3), "end_boundary": "eof"})
    return chunks


def chunk_audio_path(lid: str, idx: int) -> Path:
    return CACHE / lid / "chunks" / f"chunk_{idx:03d}.mp3"


def ensure_full_audio(L: dict) -> Path:
    flac = CACHE / L["id"] / "audio.flac"
    if not flac.exists():
        flac.parent.mkdir(parents=True, exist_ok=True)
        tmp = flac.with_name("audio.tmp.flac")
        run([FFMPEG, "-y", "-v", "error", "-i", str(COURSE / L["video"]), "-vn", "-ac", "1", "-ar", "16000",
             "-c:a", "flac", str(tmp)])
        os.replace(tmp, flac)
    return flac


def ensure_chunk_audio(L: dict, ch: dict) -> Path:
    out = chunk_audio_path(L["id"], ch["idx"])
    if out.exists() and out.stat().st_size > 0:
        return out
    flac = ensure_full_audio(L)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + ".tmp.mp3")
    run([FFMPEG, "-y", "-v", "error", "-ss", f"{ch['start']:.3f}", "-t", f"{ch['end'] - ch['start']:.3f}",
         "-i", str(flac), "-c:a", "libmp3lame", "-b:a", "48k", "-ac", "1", "-ar", "16000", str(tmp)])
    os.replace(tmp, out)
    return out


def prepare_lesson(L: dict) -> None:
    lid = L["id"]
    st = load_state(lid)
    prep = st.get("prepare")
    if not (prep and prep.get("done")):
        flac = ensure_full_audio(L)
        duration = probe_duration(flac)
        mv = mean_volume(flac)
        noise = max(-50.0, min(-28.0, mv - 14.0))
        sil = detect_silences(flac, noise, CFG["silence_min_s"], duration)
        chunks = plan_chunks(duration, sil, CFG["chunk_target_s"], CFG["chunk_window_s"], flac)
        prep = {"done": True, "at": now_iso(), "audio_duration": round(duration, 3), "mean_volume_db": mv,
                "silence_noise_db": round(noise, 1), "silence_min_s": CFG["silence_min_s"],
                "silences": sil, "chunks": chunks,
                "speech_seconds": round(duration - intervals_measure(sil), 1)}
        update_state(lid, lambda s: s.__setitem__("prepare", prep))
        log(f"prepare {lid}: {duration / 60:.1f} min, {len(chunks)} chunks, mean vol {mv:.1f} dB, "
            f"{len(sil)} silences, hard cuts={sum(c['end_boundary'] == 'hard' for c in chunks)}")
    for ch in prep["chunks"]:
        ensure_chunk_audio(L, ch)


def stage_prepare(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    with cf.ThreadPoolExecutor(max_workers=min(3, args.workers)) as ex:
        futs = {ex.submit(prepare_lesson, L): L for L in lessons}
        for f in cf.as_completed(futs):
            try:
                f.result()
            except Exception as e:  # keep going; report at the end
                log(f"prepare FAILED {futs[f]['id']}: {e}")


# --------------------------------------------------------------------------- transcribe

VOCAB = ("DTC, e-commerce, Shopify, Meta, Facebook Ads, Ads Manager, Business Manager, pixel, dataset, CAPI, "
         "ROAS, MER, CPA, CPM, CPC, CTR, AOV, LTV, COGS, CBO, ABO, ASC, Advantage+, Andromeda, UGC, hook, angle, "
         "concept, avatar, sub-avatar, core avatar, market sophistication, mechanism, desire, desire calendar, "
         "offer, bundle, BOGO, upsell, landing page, advertorial, listicle, Kalodata, Pagedeck, Toolsuite, "
         "Matedropshipping, ecomtalent, TikTok Shop, Eugene Schwartz, Breakthrough Advertising, Evolve, Origins, "
         "Skool, scaling, creative testing, feedback loop, whitebelt")

TRANSCRIBE_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "segments": {"type": "ARRAY", "items": {
            "type": "OBJECT",
            "properties": {"start": {"type": "STRING"}, "end": {"type": "STRING"},
                           "speaker": {"type": "STRING"}, "text": {"type": "STRING"}},
            "required": ["start", "end", "speaker", "text"],
            "propertyOrdering": ["start", "end", "speaker", "text"]}},
        "non_speech": {"type": "ARRAY", "items": {
            "type": "OBJECT",
            "properties": {"start": {"type": "STRING"}, "end": {"type": "STRING"}, "label": {"type": "STRING"}},
            "required": ["start", "end", "label"],
            "propertyOrdering": ["start", "end", "label"]}},
        "notes": {"type": "STRING"},
    },
    "required": ["segments", "non_speech", "notes"],
    "propertyOrdering": ["segments", "non_speech", "notes"],
}


def transcribe_prompt(L: dict, ch: dict, n: int) -> str:
    dur = ch["end"] - ch["start"]
    return f"""You are a professional transcriptionist. Produce a complete, verbatim transcript of the attached audio clip.

Context: the clip is part {ch['idx'] + 1} of {n} of the lesson "{L['title']}" from an e-commerce / DTC brand-building course ("{CFG['course_name']}" by Evolve). The clip is {dur:.1f} seconds long.

Vocabulary that may occur (use these spellings only when that is what is said): {VOCAB}.

Rules:
1. Transcribe EVERY spoken word from the first to the last second of the clip. Do not summarize, paraphrase, skip, reorder, translate or tidy up the speaker's wording. You may drop pure filler sounds ("um", "uh"); keep everything else, including numbers, names, asides and repeated phrases.
2. Split speech into segments of one sentence or short clause (about 3-15 seconds, at most ~35 words).
3. Timestamps are relative to the start of THIS clip (00:00.0 is the first moment of the clip, {hms(dur)[3:]}.0 is its end). Format them as MM:SS.s (for example "03:07.4"). "start" is when the first word of the segment begins, "end" is when the last word ends.
4. Segments must be chronological and must not overlap. Never create segments for silence.
5. Speaker: use "Presenter" for the main instructor. Use "Speaker 2", "Speaker 3"... for other people (students, guests, call participants), or their name if the audio clearly states it. Speech from a video or ad being played back uses "Played audio".
6. Unintelligible words: write [inaudible]. A word you are unsure of: write your best guess followed by [?]. Never invent content.
7. If the clip starts or ends mid-sentence, transcribe the partial words that are present.
8. List meaningful non-speech audio (music, an ad playing, long pauses, laughter) in non_speech with timestamps and a short label.
9. In notes, briefly mention anything that limited accuracy (e.g. crosstalk, music, poor audio). Use "" if nothing did.
"""


def speech_intervals(prep: dict, ch: dict) -> list[tuple[float, float]]:
    s0, s1 = ch["start"], ch["end"]
    sil = [(max(a, s0) - s0, min(b, s1) - s0) for a, b in prep["silences"] if b > s0 and a < s1]
    return intervals_subtract([(0.0, s1 - s0)], sil)


def normalize_segments(raw_segs: list, dur: float, chunk_start: float) -> tuple[list[dict], dict]:
    segs, flags = [], {"unparsed_ts": 0, "out_of_range": 0, "reordered": 0, "abs_offset_fixed": False}
    parsed = []
    for s in raw_segs or []:
        text = re.sub(r"\s+", " ", str(s.get("text", ""))).strip()
        if not text:
            continue
        a, b = parse_ts(s.get("start")), parse_ts(s.get("end"))
        parsed.append([a, b, (s.get("speaker") or "Presenter").strip() or "Presenter", text])
    starts = [p[0] for p in parsed if p[0] is not None]
    # model sometimes uses the original-video timeline; detect and undo
    if starts and chunk_start > 0 and sum(x > dur + 2 for x in starts) > 0.5 * len(starts):
        if all(-2 <= x - chunk_start <= dur + 2 for x in starts):
            for p in parsed:
                p[0] = None if p[0] is None else p[0] - chunk_start
                p[1] = None if p[1] is None else p[1] - chunk_start
            flags["abs_offset_fixed"] = True
    prev_end = 0.0
    for i, (a, b, spk, text) in enumerate(parsed):
        est = max(1.0, len(text.split()) / 2.8)
        if a is None:
            flags["unparsed_ts"] += 1
            a = prev_end
        if a > dur + 1.5:
            flags["out_of_range"] += 1
        a = min(max(a, 0.0), dur)
        if a < prev_end - 0.5:
            flags["reordered"] += 1
        a = max(a, prev_end) if a < prev_end else a
        if b is None or b <= a:
            b = a + est
        b = min(b, dur)
        if b <= a:
            b = min(dur, a + 0.5)
        segs.append({"start": round(a, 2), "end": round(b, 2), "speaker": spk, "text": text})
        prev_end = b
    # trim overlaps against the next start
    for i in range(len(segs) - 1):
        if segs[i]["end"] > segs[i + 1]["start"]:
            segs[i]["end"] = max(segs[i]["start"], segs[i + 1]["start"])
    return segs, flags


def assess_chunk(segs: list[dict], speech: list[tuple[float, float]], dur: float, flags: dict, finish: str) -> dict:
    speech_t = intervals_measure(speech)
    covered = intervals_union([(s["start"] - 0.75, s["end"] + 0.75) for s in segs])
    uncovered = intervals_subtract(speech, covered)
    # merge small separations, report long stretches of speech-like audio with no transcript
    merged = []
    for a, b in uncovered:
        if merged and a - merged[-1][1] < 1.5:
            merged[-1] = (merged[-1][0], b)
        else:
            merged.append((a, b))
    gaps = [(round(a, 1), round(b, 1)) for a, b in merged if b - a >= 6.0]
    coverage = 1.0 - (intervals_measure(uncovered) / speech_t) if speech_t > 1 else 1.0
    words = sum(len(s["text"].split()) for s in segs)
    wpm = words / (speech_t / 60) if speech_t > 1 else 0.0
    # model looping repeats whole sentences; short interjections ("Okay.", "Yeah.") legitimately repeat
    texts = [re.sub(r"\W+", "", s["text"].lower()) if len(s["text"].split()) >= 4 else "" for s in segs]
    run_len, max_run = 1, 1
    for i in range(1, len(texts)):
        run_len = run_len + 1 if texts[i] and texts[i] == texts[i - 1] else 1
        max_run = max(max_run, run_len)
    oor = flags["out_of_range"] / max(1, len(segs))
    issues = []
    if finish != "STOP":
        issues.append(f"finish_{finish}")
    if speech_t >= 20 and coverage < 0.80:
        issues.append("low_coverage")
    if speech_t >= 45 and wpm < 70:
        issues.append("low_wpm")
    if max_run >= 3:
        issues.append("repetition")
    if oor > 0.1:
        issues.append("timestamps_out_of_range")
    severe = {"low_wpm", "repetition", "timestamps_out_of_range", "finish_MAX_TOKENS"}
    passed = not (set(issues) & severe) and not (speech_t >= 20 and coverage < 0.6) and finish in ("STOP", "MAX_TOKENS") \
        and not any(i.startswith("finish_") and i != "finish_MAX_TOKENS" for i in issues)
    score = coverage - 0.4 * ("low_wpm" in issues) - 1.0 * ("repetition" in issues) - 0.5 * oor
    return {"speech_seconds": round(speech_t, 1), "coverage": round(coverage, 3), "words": words,
            "wpm": round(wpm, 1), "max_repeat_run": max_run, "gaps": gaps, "issues": issues,
            "passed": passed, "score": round(score, 3), "segments": len(segs)}


def transcribe_chunk(L: dict, ch: dict, n: int, prep: dict) -> dict:
    lid, idx = L["id"], ch["idx"]
    audio = ensure_chunk_audio(L, ch)
    b64 = base64.b64encode(audio.read_bytes()).decode("ascii")
    dur = ch["end"] - ch["start"]
    speech = speech_intervals(prep, ch)
    prompt = transcribe_prompt(L, ch, n)
    primary, fallback = CFG["transcribe_model"], CFG["fallback_model"]
    plan = [primary] * CFG["max_attempts"] + [fallback, fallback]
    attempts, best = [], None
    rawdir = RAW / lid / "transcribe"
    for k, model in enumerate(plan, 1):
        if model == fallback and best is not None:
            break  # a usable primary result exists; fallback is only for blocked/unusable chunks
        gen = {"responseMimeType": "application/json", "responseSchema": TRANSCRIBE_SCHEMA,
               "maxOutputTokens": 32768}
        if model.startswith("gemini-3"):
            gen["thinkingConfig"] = {"thinkingLevel": CFG["thinking_level"]}
        raw_path = rawdir / f"chunk_{idx:03d}.attempt{k}.json"
        resp = gemini_generate(model, [{"inlineData": {"mimeType": "audio/mp3", "data": b64}}, {"text": prompt}],
                               gen, raw_path, {"stage": "transcribe", "lesson": lid, "chunk": idx,
                                               "ref": f"{lid}#{idx}", "attempt": k})
        text, finish = response_text(resp)
        obj = loads_lenient(text) if text else None
        if not isinstance(obj, dict) or not isinstance(obj.get("segments"), list):
            attempts.append({"attempt": k, "model": model, "finish": finish, "usable": False,
                             "raw": str(raw_path.relative_to(HERE))})
            log(f"  {lid}#{idx} attempt {k} ({model}) unusable: finish={finish}")
            continue
        segs, flags = normalize_segments(obj["segments"], dur, ch["start"])
        q = assess_chunk(segs, speech, dur, flags, finish)
        non_speech = []
        for e in obj.get("non_speech") or []:
            a, b = parse_ts(e.get("start")), parse_ts(e.get("end"))
            if a is not None:
                non_speech.append({"start": round(min(max(a, 0), dur), 2),
                                   "end": round(min(max(b if b is not None else a, a), dur), 2),
                                   "label": str(e.get("label", ""))})
        result = {"model": model, "attempt": k, "finish": finish, "flags": flags, "quality": q,
                  "segments": segs, "non_speech": non_speech, "notes": obj.get("notes", ""),
                  "raw": str(raw_path.relative_to(HERE))}
        attempts.append({"attempt": k, "model": model, "finish": finish, "usable": True,
                         "passed": q["passed"], "issues": q["issues"], "coverage": q["coverage"],
                         "wpm": q["wpm"], "raw": result["raw"]})
        if best is None or q["score"] > best["quality"]["score"]:
            best = result
        if q["passed"]:
            break
        log(f"  {lid}#{idx} attempt {k} flagged {q['issues']} cov={q['coverage']} wpm={q['wpm']}; retrying")
    out = {"lesson": lid, "chunk": ch, "attempts": attempts, "completed_at": now_iso()}
    if best is None:
        out["status"] = "failed"
    else:
        out.update(best)
        out["status"] = "done" if best["quality"]["passed"] else "uncertain"
    write_json(lesson_dir_state(lid) / "chunks" / f"chunk_{idx:03d}.json", out)

    def upd(s):
        s.setdefault("transcribe", {})[str(idx)] = {"status": out["status"], "attempts": len(attempts),
                                                    "model": out.get("model"), "at": out["completed_at"]}
    update_state(lid, upd)
    q = out.get("quality", {})
    log(f"transcribe {lid}#{idx}: {out['status']} ({len(attempts)} attempt(s)) cov={q.get('coverage')} "
        f"wpm={q.get('wpm')} segs={q.get('segments')} issues={q.get('issues')}")
    return out


def stage_transcribe(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    tasks = []
    for L in lessons:
        prepare_lesson(L)
        st = load_state(L["id"])
        prep = st["prepare"]
        done = st.get("transcribe", {})
        for ch in prep["chunks"]:
            d = done.get(str(ch["idx"]), {})
            cfile = lesson_dir_state(L["id"]) / "chunks" / f"chunk_{ch['idx']:03d}.json"
            if d.get("status") in ("done", "uncertain") and cfile.exists() and not args.redo:
                continue
            if d.get("status") == "failed" and not (args.retry_failed or args.redo):
                continue
            tasks.append((L, ch, len(prep["chunks"]), prep))
    log(f"transcribe: {len(tasks)} chunk(s) to process across {len(lessons)} lesson(s); "
        f"spend so far ${usage_total():.2f}")
    stop = threading.Event()

    def work(t):
        if stop.is_set():
            return None
        try:
            return transcribe_chunk(*t)
        except (BudgetExceeded, QuotaExhausted) as e:
            stop.set()
            log(f"STOPPING: {e}")
        except Exception as e:
            log(f"transcribe ERROR {t[0]['id']}#{t[1]['idx']}: {e}\n{traceback.format_exc(limit=3)}")
            def upd(s):
                s.setdefault("transcribe", {})[str(t[1]["idx"])] = {"status": "failed", "error": str(e)[:500],
                                                                     "at": now_iso()}
            update_state(t[0]["id"], upd)
        return None

    with cf.ThreadPoolExecutor(max_workers=args.workers) as ex:
        list(ex.map(work, tasks))
    log(f"transcribe finished; total recorded spend ${usage_total():.2f}")


# --------------------------------------------------------------------------- assemble

def lesson_out(L: dict) -> Path:
    return OUT_LESSONS / L["id"]


_CONTRACT = {"gonna": ["going", "to"], "wanna": ["want", "to"], "gotta": ["got", "to"], "kinda": ["kind", "of"],
             "sorta": ["sort", "of"], "ok": ["okay"], "cause": ["because"], "cuz": ["because"], "yeah": ["yes"],
             "yep": ["yes"], "alright": ["all", "right"]}
_NUMWORDS = {"zero": "0", "one": "1", "two": "2", "three": "3", "four": "4", "five": "5", "six": "6", "seven": "7",
             "eight": "8", "nine": "9", "ten": "10", "eleven": "11", "twelve": "12", "fifteen": "15", "twenty": "20",
             "thirty": "30", "forty": "40", "fifty": "50", "hundred": "100", "thousand": "000"}


def norm_tokens(text: str) -> list[str]:
    # letters and digit runs as separate tokens, so "2.2x", "2 .2 X" and "$10,000" vs "$10 ,000" compare equal;
    # contractions and number words normalised so "gonna"/"going to" and "two X"/"2x" match
    out = []
    for t in re.findall(r"[a-z]+|\d+", text.lower()):
        out.extend(_CONTRACT.get(t) or [_NUMWORDS.get(t, t)])
    return out


_NUMBERISH = {"a", "and", "point", "dollars", "dollar", "bucks", "percent", "k", "x", "grand", "million",
              "thirteen", "fourteen", "sixteen", "seventeen", "eighteen", "nineteen", "sixty", "seventy",
              "eighty", "ninety", "m", "s"}


def _join_words(words: list[str]) -> str:
    s = " ".join(w.strip() for w in words if w.strip())
    s = re.sub(r"(\d) ([.,]\d)", r"\1\2", s)
    return re.sub(r"\s+([,.?!%])", r"\1", s)


_STOP = set("a an the and or but so to of in on at for with from is are was were be been being it its that this "
            "these those i you he she we they me my your our their them us him her do does did have has had not no "
            "yes okay like just if then there here what how why when going can will would should could get got "
            "know really um uh all right s t re ll ve d m because as about up out one 1".split())


def content_ratio(tokens: list[str], heard: list[str]) -> float:
    """Share of the content words in `tokens` that also occur in `heard`."""
    from collections import Counter
    want = Counter(t for t in tokens if t not in _STOP)
    have = Counter(heard)
    if not want:
        return 1.0
    return sum(min(n, have[t]) for t, n in want.items()) / sum(want.values())


def align_timing(segs: list[dict], wwords: list, chunks: list[dict], checks: dict | None = None):
    """Re-time Gemini segments with local Whisper word timestamps, and cross-check completeness.

    Gemini's text is kept; its own timestamps drift by up to ~30 s inside a 5-minute chunk. Gemini words are
    aligned to Whisper words chunk by chunk (difflib, monotonic); segment start/end come from the matched
    Whisper words and unmatched segments are interpolated between aligned neighbours. Disagreements are then
    verified against nearby words: speech Whisper heard that is missing from the Gemini text is inserted as a
    marked fill segment, and Gemini text Whisper cannot confirm is flagged.
    """
    import difflib
    from collections import Counter

    ww = []  # (token, start, end, whisper word index)
    for wi, w in enumerate(wwords):
        for t in norm_tokens(w[2]):
            ww.append((t, w[0], w[1], wi))
    total_g = total_m = 0
    seg_tok = [norm_tokens(s["text"]) for s in segs]
    match_times: list[list[tuple[int, float, float]]] = [[] for _ in segs]
    cands = []
    wp = 0  # whisper pointer: alignment is sequential and monotonic across chunks
    for ch in chunks:
        idxs = [i for i, s in enumerate(segs) if s["chunk"] == ch["idx"]]
        g = [(t, i, k) for i in idxs for k, t in enumerate(seg_tok[i])]
        if not g:
            continue
        j0 = wp
        while j0 > 0 and ww[j0 - 1][1] >= ch["start"] - 3:
            j0 -= 1
        j1 = j0
        while j1 < len(ww) and ww[j1][1] <= ch["end"] + 3:
            j1 += 1
        win = ww[j0:j1]
        sm = difflib.SequenceMatcher(a=[x[0] for x in g], b=[x[0] for x in win], autojunk=False)
        matched = 0
        for a, b, n in sm.get_matching_blocks():
            if n < 2:  # single-word matches are too often coincidental ("the", "and")
                continue
            for k in range(n):
                _, si, pos = g[a + k]
                _, ws, we, _ = win[b + k]
                match_times[si].append((pos, ws, we))
            matched += n
            wp = max(wp, j0 + b + n)
        total_g += len(g)
        total_m += matched
        for tag, a0, a1, b0, b1 in sm.get_opcodes():
            na, nb = a1 - a0, b1 - b0
            if tag in ("insert", "replace") and nb >= 6 and na < 0.3 * nb:
                cands.append(("omission", [j0 + k for k in range(b0, b1)], ch["idx"]))
            if tag in ("delete", "replace") and na >= 12 and nb < 0.3 * na:
                cands.append(("unverified", [g[k] for k in range(a0, a1)], ch["idx"]))
    out = []
    for i, s in enumerate(segs):
        s2 = dict(s, gemini_start=s["start"], gemini_end=s["end"])
        mt = match_times[i]
        n = max(1, len(seg_tok[i]))
        if mt:
            mt.sort()
            p0, ws, _ = mt[0]
            p1, ws1, we = mt[-1]
            s2["start"] = round(max(0.0, ws - 0.3 * p0), 2)
            s2["end"] = round(we + 0.3 * (n - 1 - p1), 2)
            s2["_spoken_until"] = ws1 + 0.35 + 0.3 * (n - 1 - p1)  # word END times can stretch over pauses
            s2["timing"] = "aligned" if len(mt) >= 0.5 * n else "partly_aligned"
        else:
            s2["timing"] = "interpolated"
        out.append(s2)
    # Text the audio may not contain: a run of segments base.en did not hear, squeezed between two aligned
    # neighbours with no time in which it could have been spoken (> 6 words/s). On lesson 1 Gemini inserted
    # ~56 plausible words inside 1.5 s. base.en also skips real speech in crosstalk, so a run is only a
    # candidate: the window is re-transcribed with small.en (beam search) and the text is removed only if
    # small.en does not hear its content words either. Candidates not yet checked are kept and flagged.
    removed_runs, pending, check_issues = [], [], []
    i = 0
    while i < len(out):
        if out[i]["timing"] == "aligned":
            i += 1
            continue
        j = i
        while j < len(out) and out[j]["timing"] != "aligned":
            j += 1
        if i > 0 and j < len(out):
            avail = out[j]["start"] - out[i - 1].get("_spoken_until", out[i - 1]["end"])
            words = sum(len(seg_tok[k]) for k in range(i, j))
            if words >= 8 and words > 6.0 * max(avail, 0.25):
                text = " ".join(out[k]["text"] for k in range(i, j))
                key = hashlib.sha1(f"{out[i]['chunk']}|{text}".encode("utf-8")).hexdigest()[:12]
                res = (checks or {}).get(key)
                run_tokens = [t for k in range(i, j) for t in seg_tok[k]]
                if res is None:
                    pending.append({"key": key, "start": round(max(0.0, out[i - 1]["start"] - 1.0), 2),
                                    "end": round(out[j]["end"] + 1.0, 2), "text": text})
                    check_issues.append({"type": "pending_audio_check", "start": round(out[i - 1]["start"], 1),
                                         "end": round(out[j]["start"], 1), "chunk": out[i]["chunk"],
                                         "detail": f"{words} word(s) the fast check did not hear; awaiting small.en "
                                                   f"re-check: \"{text[:150]}\""})
                elif content_ratio(run_tokens, norm_tokens(res["text"])) < 0.35:
                    for k in range(i, j):
                        out[k]["timing"] = "not_in_audio"
                    removed_runs.append((out[i - 1]["start"], out[j]["start"], list(range(i, j))))
        i = j
    keep = [k for k in range(len(out)) if out[k]["timing"] != "not_in_audio"]
    newidx = {old: new for new, old in enumerate(keep)}
    removed = [dict(out[k], removed_reason="not heard by local Whisper base.en nor by a small.en re-check of "
                                           "the window, and no time to have been spoken between the "
                                           "neighbouring aligned segments")
               for k in range(len(out)) if out[k]["timing"] == "not_in_audio"]
    removal_issues = [{"type": "removed_unspoken_text", "start": round(a, 1), "end": round(max(a, b), 1),
                       "chunk": out[ks[0]]["chunk"],
                       "detail": f"{sum(len(seg_tok[k]) for k in ks)} word(s) in the Gemini text are not in the audio "
                                 f"and were removed (kept in transcript.json \"removed_segments\"): "
                                 f"\"{' '.join(out[k]['text'] for k in ks)[:200]}\""}
                      for a, b, ks in removed_runs]
    out = [out[k] for k in keep]
    seg_tok = [seg_tok[k] for k in keep]
    remapped = []
    for kind, items, chidx in cands:
        if kind == "unverified":
            items = [(t, newidx[si], pos) for t, si, pos in items if si in newidx]
            if len(items) < 12:
                continue
        remapped.append((kind, items, chidx))
    cands = remapped
    # interpolate unaligned runs between aligned neighbours (by word count); edges keep a local Gemini offset
    i = 0
    while i < len(out):
        if out[i]["timing"] != "interpolated":
            i += 1
            continue
        j = i
        while j < len(out) and out[j]["timing"] == "interpolated":
            j += 1
        prev = out[i - 1] if i > 0 else None
        nxt = out[j] if j < len(out) else None
        if prev and nxt:
            t0, t1 = prev["end"], max(prev["end"], nxt["start"])
            words = [max(1, len(seg_tok[k])) for k in range(i, j)]
            tot = sum(words)
            cur = t0
            for k, wcount in zip(range(i, j), words):
                d = (t1 - t0) * wcount / tot
                out[k]["start"], out[k]["end"] = round(cur, 2), round(cur + d, 2)
                cur += d
        else:
            ref = prev or nxt
            off = (ref["start"] - ref["gemini_start"]) if ref else 0.0
            for k in range(i, j):
                out[k]["start"] = round(max(0.0, out[k]["gemini_start"] + off), 2)
                out[k]["end"] = round(max(out[k]["start"], out[k]["gemini_end"] + off), 2)
        i = j

    def gem_bag(t0, t1):
        bag = Counter()
        for k, s in enumerate(out):
            if s["end"] >= t0 and s["start"] <= t1:
                bag.update(seg_tok[k])
        return bag

    def found_ratio(toks, bag):
        c = Counter(toks)
        return sum(min(n, bag[t]) for t, n in c.items()) / max(1, sum(c.values()))

    issues, fills, filled = list(removal_issues) + check_issues, [], set()
    for kind, items, chidx in cands:
        if kind == "omission":
            idx = [j for j in items if ww[j][3] not in filled]
            if len(idx) < 6:
                continue
            numeric = sum(1 for j in idx if ww[j][0].isdigit() or ww[j][0] in _NUMBERISH)
            if numeric >= 0.5 * len(idx):
                continue  # "a hundred and twenty six dollars" vs "$126": same words, different notation
            t0, t1 = ww[idx[0]][1], ww[idx[-1]][2]
            if found_ratio([ww[j][0] for j in idx], gem_bag(t0 - 15, t1 + 15)) >= 0.6:
                continue  # the words are there, just aligned differently
            wis = sorted({ww[j][3] for j in idx})
            text = _join_words([wwords[k][2] for k in wis])
            probs = [wwords[k][3] for k in wis if len(wwords[k]) > 3]
            if probs and statistics.mean(probs) < 0.5:
                issues.append({"type": "possible_omission", "start": round(t0, 1), "end": round(t1, 1),
                               "chunk": chidx, "detail": f"local Whisper (low confidence) heard ~{len(wis)} words "
                                                         f"missing from the transcript: \"{text[:150]}\""})
                continue
            filled.update(wis)
            fills.append({"start": round(t0, 2), "end": round(t1, 2), "speaker": None, "text": text,
                          "chunk": chidx, "timing": "whisper", "source": "local_whisper_fill",
                          "gemini_start": None, "gemini_end": None})
            issues.append({"type": "filled_omission", "start": round(t0, 1), "end": round(t1, 1), "chunk": chidx,
                           "detail": f"{len(wis)} word(s) heard by local Whisper but missing from the Gemini text "
                                     f"were inserted, marked [+ ...]: \"{text[:150]}\""})
        else:
            si0, si1 = items[0][1], items[-1][1]
            t0, t1 = out[si0]["start"], out[si1]["end"]
            bag = Counter(x[0] for x in ww if t0 - 15 <= x[1] <= t1 + 15)
            if found_ratio([x[0] for x in items], bag) >= 0.6:
                continue
            ptext = " ".join(out[k]["text"] for k in range(si0, si1 + 1))
            key = hashlib.sha1(f"confirm|{chidx}|{ptext}".encode("utf-8")).hexdigest()[:12]
            res = (checks or {}).get(key)
            if res is None:
                pending.append({"key": key, "start": round(max(0.0, t0 - 4), 2), "end": round(t1 + 4, 2),
                                "text": ptext})
                detail = f"~{len(items)} transcript words not heard by local Whisper base.en; awaiting small.en re-check"
            elif content_ratio([x[0] for x in items], norm_tokens(res["text"])) >= 0.5:
                continue  # the stronger check hears it: base.en simply skipped this speech
            else:
                detail = (f"~{len(items)} transcript words heard by neither local Whisper base.en nor small.en "
                          "(music, played audio, foreign speech, crosstalk, or a transcription error)")
            issues.append({"type": "unverified_passage", "start": round(t0, 1), "end": round(t1, 1), "chunk": chidx,
                           "detail": detail})
    for f in fills:  # speaker of the segment just before the fill
        before = [s for s in out if s["start"] <= f["start"]]
        f["speaker"] = before[-1]["speaker"] if before else (out[0]["speaker"] if out else "Presenter")
    out = sorted(out + fills, key=lambda s: s["start"])
    for k in range(1, len(out)):  # enforce order and no overlaps
        if out[k]["start"] < out[k - 1]["start"]:
            out[k]["start"] = out[k - 1]["start"]
        if out[k - 1]["end"] > out[k]["start"]:
            out[k - 1]["end"] = out[k]["start"]
        if out[k]["end"] < out[k]["start"]:
            out[k]["end"] = out[k]["start"]
    stats = {"gemini_tokens": total_g, "matched_tokens": total_m,
             "match_ratio": round(total_m / total_g, 3) if total_g else 0.0,
             "segments_aligned": sum(1 for s in out if s["timing"] == "aligned"),
             "segments_partly_aligned": sum(1 for s in out if s["timing"] == "partly_aligned"),
             "segments_interpolated": sum(1 for s in out if s["timing"] == "interpolated"),
             "whisper_fills": len(fills), "whisper_fill_words": len(filled),
             "removed_segments": len(removed), "removed_words": sum(len(norm_tokens(s["text"])) for s in removed),
             "audio_checks_pending": len(pending)}
    return out, stats, issues, removed, pending


def stage_assemble(args) -> None:
    for L in select_lessons(args.lessons, video_only=True):
        assemble_lesson(L)


def assemble_lesson(L: dict) -> dict | None:
    lid = L["id"]
    st = load_state(lid)
    prep = st.get("prepare")
    if not prep:
        log(f"assemble {lid}: not prepared; skipping")
        return None
    chunks_meta, segs, non_speech, issues = [], [], [], []
    wpms = []
    for ch in prep["chunks"]:
        c = read_json(lesson_dir_state(lid) / "chunks" / f"chunk_{ch['idx']:03d}.json")
        status = c["status"] if c else "pending"
        meta = {"idx": ch["idx"], "start": ch["start"], "end": ch["end"], "end_boundary": ch["end_boundary"],
                "status": status}
        if c and c.get("segments") is not None:
            q = c["quality"]
            meta.update({"model": c["model"], "attempts": len(c["attempts"]), "coverage": q["coverage"],
                         "wpm": q["wpm"], "issues": q["issues"], "notes": c.get("notes", "")})
            if q["speech_seconds"] >= 45:
                wpms.append(q["wpm"])
            for s in c["segments"]:
                segs.append({"start": round(ch["start"] + s["start"], 2), "end": round(ch["start"] + s["end"], 2),
                             "speaker": s["speaker"], "text": s["text"], "chunk": ch["idx"]})
            for e in c.get("non_speech", []):
                non_speech.append({"start": round(ch["start"] + e["start"], 2),
                                   "end": round(ch["start"] + e["end"], 2), "label": e["label"]})
            for a, b in q["gaps"]:
                issues.append({"type": "possible_missing_speech", "start": round(ch["start"] + a, 1),
                               "end": round(ch["start"] + b, 1), "chunk": ch["idx"],
                               "detail": "audio with speech-level energy but no transcript segment"})
            if status == "uncertain":
                issues.append({"type": "uncertain_chunk", "start": ch["start"], "end": ch["end"], "chunk": ch["idx"],
                               "detail": f"best attempt still flagged: {', '.join(q['issues'])}"})
            if c.get("notes"):
                issues.append({"type": "transcriber_note", "start": ch["start"], "end": ch["end"],
                               "chunk": ch["idx"], "detail": c["notes"]})
        else:
            issues.append({"type": "missing_chunk" if status == "failed" else "not_transcribed_yet",
                           "start": ch["start"], "end": ch["end"], "chunk": ch["idx"],
                           "detail": "no usable transcription" if status == "failed" else "pending"})
        if ch["end_boundary"] == "hard":
            issues.append({"type": "hard_chunk_boundary", "start": ch["end"], "end": ch["end"], "chunk": ch["idx"],
                           "detail": "no pause found near the cut; a word may be split across chunks"})
        chunks_meta.append(meta)
    align_stats, removed = None, []
    wfile = lesson_dir_state(lid) / CFG.get("whisper_file", "whisper_words.json")
    if wfile.exists() and segs:
        wdoc = read_json(wfile)
        checks = read_json(lesson_dir_state(lid) / "audio_checks.json", {}) or {}
        segs, align_stats, a_issues, removed, pending = align_timing(segs, wdoc["words"], prep["chunks"], checks)
        write_json(lesson_dir_state(lid) / "audio_checks_pending.json", pending)
        align_stats["whisper_model"] = wdoc["model"]
        # gap checks based on Gemini's drifting timestamps are superseded by the word-level comparison
        issues = [i for i in issues if i["type"] != "possible_missing_speech"] + a_issues
    if wpms:
        med = statistics.median(wpms)
        for m in chunks_meta:
            if m.get("wpm") and m["wpm"] < 0.55 * med and "low_wpm" not in (m.get("issues") or []):
                issues.append({"type": "low_word_rate", "start": m["start"], "end": m["end"], "chunk": m["idx"],
                               "detail": f"{m['wpm']} wpm vs lesson median {med:.0f}; possible omissions"})
    for i, s in enumerate(segs, 1):
        s["id"] = i
        if "[inaudible]" in s["text"] or "[?]" in s["text"]:
            issues.append({"type": "unclear_words", "start": s["start"], "end": s["end"], "chunk": s["chunk"],
                           "detail": s["text"][:160]})
    issues.sort(key=lambda x: (x["start"], x["type"]))
    complete = all(m["status"] in ("done", "uncertain") for m in chunks_meta)
    models = sorted({m.get("model") for m in chunks_meta if m.get("model")})
    speakers = sorted({s["speaker"] for s in segs})
    doc = {
        "lesson": {k: L.get(k) for k in ("num", "id", "title", "section", "video", "duration", "skool_url")},
        "generated_at": now_iso(),
        "status": "complete" if complete else "partial",
        "timeline": "All times are seconds on the original video timeline.",
        "backend": {"engine": "Google Gemini API", "models": models,
                    "method": f"audio cut at pauses into ~{CFG['chunk_target_s']}s chunks, chunk offsets added back",
                    "style": "verbatim; pure filler sounds (um/uh) may be omitted",
                    "timing": (f"segment times from local faster-whisper {align_stats['whisper_model']} word "
                               f"timestamps aligned to the Gemini text ({align_stats['match_ratio']:.0%} of words "
                               "matched); unmatched segments interpolated") if align_stats else
                              "Gemini's own timestamps (NOT yet aligned; can drift up to ~30 s within a chunk)"},
        "alignment": align_stats,
        "speakers": speakers,
        "chunks": chunks_meta,
        "issues": issues,
        "segments": [{k: s[k] for k in ("id", "start", "end", "speaker", "text", "chunk", "timing", "source")
                      if k in s} for s in segs],
        "removed_segments": [{k: s.get(k) for k in ("speaker", "text", "chunk", "gemini_start", "gemini_end",
                                                    "removed_reason")} for s in removed],
        "non_speech": non_speech,
    }
    out = lesson_out(L)
    write_json(out / "transcript.json", doc)
    write_text(out / "transcript.srt", render_srt(segs, multi=len(speakers) > 1))
    write_text(out / "transcript.md", render_transcript_md(L, doc))
    update_state(lid, lambda s: s.__setitem__("assemble", {"at": now_iso(), "status": doc["status"],
                                                           "segments": len(segs), "issues": len(issues),
                                                           "aligned": bool(align_stats)}))
    log(f"assemble {lid}: {doc['status']}, {len(segs)} segments, {len(issues)} issue(s), timing "
        + (f"aligned ({align_stats['match_ratio']:.0%} words matched)" if align_stats else "NOT aligned"))
    return doc


def render_srt(segs: list[dict], multi: bool) -> str:
    lines = []
    for i, s in enumerate(segs, 1):
        text = f"{s['speaker']}: {s['text']}" if multi else s["text"]
        if s.get("source") == "local_whisper_fill":
            text = "[+] " + text
        lines += [str(i), f"{srt_ts(s['start'])} --> {srt_ts(max(s['end'], s['start'] + 0.3))}", text, ""]
    return "\n".join(lines)


ISSUE_LABEL = {
    "possible_missing_speech": "Possible missing speech",
    "uncertain_chunk": "Uncertain section",
    "missing_chunk": "MISSING (transcription failed)",
    "not_transcribed_yet": "Not transcribed yet",
    "hard_chunk_boundary": "Chunk boundary without a pause",
    "low_word_rate": "Low word rate (possible omissions)",
    "unclear_words": "Unclear words",
    "transcriber_note": "Transcriber note",
    "possible_omission": "Possible omission",
    "filled_omission": "Filled from second transcriber",
    "removed_unspoken_text": "Removed text that is not in the audio",
    "pending_audio_check": "Awaiting second audio check",
    "unverified_passage": "Passage not confirmed by second transcriber",
}


def render_transcript_md(L: dict, doc: dict) -> str:
    segs = doc["segments"]
    multi = len(doc["speakers"]) > 1
    video_link = rel_link(COURSE / L["video"], lesson_out(L))
    head = [f"# {L['num']}. {L['title']} - Transcript", "",
            f"- **Course section:** {L.get('section', '')}",
            f"- **Source video:** [{L['video']}]({video_link}) ({hms(L.get('duration', 0))})",
            f"- **Status:** {doc['status']} - {len(segs)} segments",
            f"- **Transcribed with:** {', '.join(doc['backend']['models']) or 'n/a'} "
            f"(Gemini API, {doc['backend']['method']})",
            f"- **Timestamps:** `[HH:MM:SS]` on the original video timeline; {doc['backend']['timing']}",
            "- **Other formats:** [transcript.json](transcript.json) | [transcript.srt](transcript.srt)",
            "", "> AI-generated verbatim transcript (filler sounds may be omitted). Summaries are kept separately "
            "in `summary.md`." + (" Text shown as *[+ ...]* was missed by the primary transcriber and filled in "
                                   "from a local Whisper pass." if any(s.get("source") for s in segs) else ""), ""]
    flagged = [i for i in doc["issues"] if i["type"] not in ("transcriber_note",)]
    if flagged:
        head += ["## Flagged sections", ""]
        for i in flagged[:60]:
            rng = hms(i["start"]) if i["start"] == i["end"] else f"{hms(i['start'])}-{hms(i['end'])}"
            head.append(f"- `[{rng}]` {ISSUE_LABEL.get(i['type'], i['type'])}: {i['detail']}")
        if len(flagged) > 60:
            head.append(f"- ... {len(flagged) - 60} more in transcript.json")
        head.append("")
    notes = [i for i in doc["issues"] if i["type"] == "transcriber_note"]
    if notes:
        head += ["<details><summary>Transcriber notes</summary>", "",
                 "Notes are written per ~5-minute chunk (range shown); any times mentioned inside a note are "
                 "relative to that chunk's start.", ""]
        head += [f"- `[{hms(i['start'])}-{hms(i['end'])}]` {i['detail']}" for i in notes]
        head += ["", "</details>", ""]
    head += ["## Transcript", ""]
    gap_marks = {(i["start"]): i for i in doc["issues"]
                 if i["type"] in ("possible_missing_speech", "missing_chunk", "not_transcribed_yet",
                                  "possible_omission")}
    body, para, p_start, p_spk, p_words = [], [], None, None, 0
    pending_marks = sorted(gap_marks)

    def flush():
        nonlocal para, p_start, p_spk, p_words
        if para:
            spk = f"**{p_spk}:** " if multi else ""
            body.append(f"**[{hms(p_start)}]** {spk}{' '.join(para)}")
            body.append("")
        para, p_start, p_spk, p_words = [], None, None, 0

    for s in segs:
        while pending_marks and pending_marks[0] <= s["start"]:
            flush()
            g = gap_marks[pending_marks.pop(0)]
            body += [f"> **[{hms(g['start'])}-{hms(g['end'])}] {ISSUE_LABEL[g['type']]}.** {g['detail']}", ""]
        new_para = (p_start is None or s["speaker"] != p_spk or s["start"] - p_start > 60 or p_words > 140)
        if new_para:
            flush()
            p_start, p_spk = s["start"], s["speaker"]
        para.append(f"*[+ {s['text']}]*" if s.get("source") == "local_whisper_fill" else s["text"])
        p_words += len(s["text"].split())
    flush()
    for t in pending_marks:
        g = gap_marks[t]
        body += [f"> **[{hms(g['start'])}-{hms(g['end'])}] {ISSUE_LABEL[g['type']]}.** {g['detail']}", ""]
    return "\n".join(head + body)


# --------------------------------------------------------------------------- frames (candidates)

def stage_frames(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    with cf.ThreadPoolExecutor(max_workers=min(3, args.workers)) as ex:
        futs = {ex.submit(frames_lesson, L): L for L in lessons}
        for f in cf.as_completed(futs):
            try:
                f.result()
            except Exception as e:
                log(f"frames FAILED {futs[f]['id']}: {e}\n{traceback.format_exc(limit=3)}")


def frames_lesson(L: dict) -> None:
    import numpy as np
    from PIL import Image

    lid = L["id"]
    st = load_state(lid)
    if st.get("frames", {}).get("done"):
        return
    kdir = CACHE / lid / "keyframes"
    meta_path = kdir / "keyframes.json"
    kf = read_json(meta_path)
    if not kf:
        if kdir.exists():
            shutil.rmtree(kdir)
        kdir.mkdir(parents=True)
        p = run([FFMPEG, "-hide_banner", "-skip_frame", "nokey", "-i", str(COURSE / L["video"]), "-an",
                 "-fps_mode", "passthrough", "-vf", "showinfo", "-q:v", "2", str(kdir / "kf_%05d.jpg")])
        times = [float(m.group(1)) for m in re.finditer(r"n:\s*\d+.*?pts_time:\s*([\d.]+)", p.stderr)]
        files = sorted(kdir.glob("kf_*.jpg"))
        if len(times) != len(files):
            raise RuntimeError(f"keyframe count mismatch {len(times)} vs {len(files)}")
        kf = [{"t": round(t, 3), "file": f.name} for t, f in zip(times, files)]
        write_json(meta_path, kf)

    def feat(name):
        im = Image.open(kdir / name).convert("L").resize((64, 36), Image.BILINEAR)
        return np.asarray(im, dtype=np.float32)

    feats = [feat(k["file"]) for k in kf]
    # Split into visual scenes: a new scene starts when the frame differs enough from the scene's first frame.
    T_SCENE = 7.0
    scenes, cur = [], [0]
    for i in range(1, len(kf)):
        if float(np.abs(feats[i] - feats[cur[0]]).mean()) >= T_SCENE:
            scenes.append(cur)
            cur = [i]
        else:
            cur.append(i)
    scenes.append(cur)
    # Representative = last frame of a scene (slides are fully built just before they change);
    # long scenes also get a frame every 90 s to catch slow changes (scrolling, typing).
    picks = []
    for sc in scenes:
        idxs = [sc[-1]]
        t0 = kf[sc[0]]["t"]
        for i in sc:
            if kf[i]["t"] - t0 >= 90:
                idxs.append(i)
                t0 = kf[i]["t"]
        picks.extend(sorted(set(idxs)))
    picks = sorted(set(picks))
    # drop near-duplicates of any of the last 6 kept frames (slide <-> webcam toggling)
    kept = []
    for i in picks:
        if any(float(np.abs(feats[i] - feats[j]).mean()) < 3.0 for j in kept[-6:]):
            continue
        kept.append(i)
    cap = max(12, int(L["duration"] / 8))
    if len(kept) > cap:  # spread evenly if a lesson is extremely busy
        step = len(kept) / cap
        kept = [kept[int(k * step)] for k in range(cap)]
    cands = [{"cid": f"F{n + 1:03d}", "t": kf[i]["t"], "file": kf[i]["file"]} for n, i in enumerate(kept)]
    update_state(lid, lambda s: s.__setitem__("frames", {"done": True, "at": now_iso(), "keyframes": len(kf),
                                                         "scenes": len(scenes), "candidates": cands}))
    log(f"frames {lid}: {len(kf)} keyframes -> {len(scenes)} scenes -> {len(cands)} candidates")


# --------------------------------------------------------------------------- visual notes

VISUAL_SCHEMA = {
    "type": "OBJECT",
    "properties": {"frames": {"type": "ARRAY", "items": {
        "type": "OBJECT",
        "properties": {
            "frame_id": {"type": "STRING"},
            "category": {"type": "STRING", "enum": [
                "slide", "diagram", "spreadsheet", "chart", "ad_example", "software_demo", "webpage",
                "product_page", "document", "social_post", "talking_head", "title_card", "transition", "other"]},
            "useful": {"type": "BOOLEAN"},
            "duplicate_of": {"type": "STRING"},
            "title": {"type": "STRING"},
            "on_screen": {"type": "STRING"},
            "not_in_transcript": {"type": "STRING"},
            "legibility": {"type": "STRING", "enum": ["clear", "partly legible", "illegible", "n/a"]},
        },
        "required": ["frame_id", "category", "useful", "duplicate_of", "title", "on_screen",
                     "not_in_transcript", "legibility"],
        "propertyOrdering": ["frame_id", "category", "useful", "duplicate_of", "title", "on_screen",
                             "not_in_transcript", "legibility"]}}},
    "required": ["frames"],
}


def transcript_excerpt(doc: dict, t0: float, t1: float) -> str:
    multi = len(doc["speakers"]) > 1
    lines = []
    for s in doc["segments"]:
        if s["end"] >= t0 and s["start"] <= t1:
            spk = f"{s['speaker']}: " if multi else ""
            lines.append(f"[{hms(s['start'])}] {spk}{s['text']}")
    return "\n".join(lines) or "(no speech in this window)"


def visual_prompt(L: dict, batch: list[dict], prev_titles: list[str], excerpt: str, t0: float, t1: float) -> str:
    prev = "\n".join(f"- {p}" for p in prev_titles[-8:]) or "- (none yet)"
    return f"""You are documenting the on-screen content of a recorded lesson so that a reader who only has the transcript misses nothing important.

Lesson: "{L['title']}" (course section: {L.get('section', '')}) from an e-commerce / DTC brand-building course.
You receive {len(batch)} still frames sampled from the video (each preceded by its id and timestamp) and the spoken transcript around those times.

For EACH frame return one object:
- frame_id: the id given.
- category: the best matching category.
- useful: true if the frame carries information worth keeping: readable slide text, a diagram or framework, a spreadsheet, numbers, an ad creative, a product or landing page, social posts, or software UI being demonstrated (settings, menus, metrics). false for presenter-only shots, blank, blurred or transition frames, and for frames that are near-duplicates of a more complete frame in this batch or of a previously kept frame listed below.
- duplicate_of: id of the more complete frame in this batch showing the same content, else "".
- title: short descriptive title (max 12 words).
- on_screen: faithful, concise record of the informative content. Slides: headings and bullet text verbatim. Diagrams: structure and labels. Spreadsheets/tables: column headers, row labels and the key numbers. Ads: headline, body copy, offer, visual concept. Software: which tool and screen, what is selected or configured, visible values. "" when not useful.
- not_in_transcript: information that is visible on screen but NOT said aloud in the transcript excerpt (exact figures, labels, settings, URLs, names, copy, layout). Be specific. "" if everything shown is spoken or nothing informative is shown.
- legibility: clear / partly legible / illegible / n/a.

Rules: describe only what is visible; write [illegible] for text you cannot read; never guess. Do not summarize the lesson.
A frame partly covered by a pop-up, cursor or webcam overlay is still useful if its readable part shows information that no other kept frame shows (for example the final numbers of a calculation); record what is readable and mark it partly legible.
Copy URLs, codes and figures character by character; mark any uncertain character with [?].

Frames already kept earlier in this lesson (avoid duplicating them):
{prev}

Transcript excerpt {hms(t0)}-{hms(t1)}:
{excerpt}
"""


def stage_visual(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    stop = threading.Event()

    def work(L):
        if stop.is_set():
            return
        try:
            visual_lesson(L, redo=args.redo)
        except (BudgetExceeded, QuotaExhausted) as e:
            stop.set()
            log(f"STOPPING: {e}")
        except Exception as e:
            log(f"visual ERROR {L['id']}: {e}\n{traceback.format_exc(limit=3)}")

    with cf.ThreadPoolExecutor(max_workers=args.workers) as ex:
        list(ex.map(work, lessons))
    log(f"visual finished; total recorded spend ${usage_total():.2f}")


PARAPHRASE_RULE = """
Copyright rule for this request: some frames show long passages of published text (book pages, articles, long
reviews or posts). Do NOT transcribe such passages. Describe what the text is (title, source, topic) and summarize its
key points in your own words, quoting at most a short phrase. Slides, short UI labels, numbers and ad headlines may
still be copied exactly.
"""


def describe_frames(L: dict, doc: dict, batch: list[dict], prev_titles: list[str], raw_stem: str,
                    paraphrase: bool, attempts: int) -> tuple[list | None, str]:
    """One Gemini request for a batch of candidate frames; returns (frames, last finishReason)."""
    lid = L["id"]
    kdir = CACHE / lid / "keyframes"
    t0, t1 = max(0, batch[0]["t"] - 45), batch[-1]["t"] + 30
    parts = []
    for c in batch:
        parts.append({"text": f"Frame {c['cid']} at {hms(c['t'])}"})
        parts.append({"inlineData": {"mimeType": "image/jpeg",
                                     "data": base64.b64encode((kdir / c["file"]).read_bytes()).decode()}})
    prompt = visual_prompt(L, batch, prev_titles, transcript_excerpt(doc, t0, t1), t0, t1)
    parts.append({"text": prompt + (PARAPHRASE_RULE if paraphrase else "")})
    gen = {"responseMimeType": "application/json", "responseSchema": VISUAL_SCHEMA, "maxOutputTokens": 32768}
    model = CFG["visual_model"]
    if model.startswith("gemini-3"):
        gen["thinkingConfig"] = {"thinkingLevel": CFG["thinking_level"]}
    finish = ""
    for attempt in range(1, attempts + 1):
        raw_path = RAW / lid / "visual" / f"{raw_stem}.attempt{attempt}.json"
        resp = gemini_generate(model, parts, gen, raw_path, {"stage": "visual", "lesson": lid, "batch": raw_stem,
                                                             "ref": f"{lid}/{raw_stem}", "attempt": attempt})
        text, finish = response_text(resp)
        obj = loads_lenient(text) if text else None
        if isinstance(obj, dict) and isinstance(obj.get("frames"), list):
            return obj["frames"], finish
        log(f"  visual {lid} {raw_stem} attempt {attempt} unusable (finish={finish})")
        if finish == "RECITATION":
            break  # repeating the identical request does not help; caller switches strategy
    return None, finish


def visual_lesson(L: dict, redo: bool = False) -> None:
    lid = L["id"]
    st = load_state(lid)
    if st.get("visual", {}).get("done") and not redo and not st["visual"].get("unreviewed"):
        return
    doc = read_json(lesson_out(L) / "transcript.json")
    if not doc:
        log(f"visual {lid}: no transcript yet; skipping (transcripts are reused, never re-made here)")
        return
    if not doc.get("alignment") and not CFG.get("allow_unaligned"):
        log(f"visual {lid}: transcript timing not aligned yet (waiting for Whisper pass); skipping")
        return
    if not st.get("frames", {}).get("done"):
        frames_lesson(L)
        st = load_state(lid)
    cands = st["frames"]["candidates"]
    kdir = CACHE / lid / "keyframes"
    bsize = CFG["visual_batch"]
    vdir = lesson_dir_state(lid) / "visual"
    results: dict[str, dict] = {}
    prev_titles: list[str] = []
    for bi in range(0, len(cands), bsize):
        batch = cands[bi:bi + bsize]
        bpath = vdir / f"batch_{bi // bsize:03d}.json"
        saved = read_json(bpath)
        reviewed_ids = {f.get("frame_id") for f in (saved or {}).get("frames", [])}
        if saved and not redo and all(c["cid"] in reviewed_ids for c in batch):
            frames = saved["frames"]
        else:
            if any(not (kdir / c["file"]).exists() for c in batch):
                update_state(lid, lambda s: s.pop("frames", None))
                frames_lesson(L)  # cache was cleared; rebuild (same deterministic candidates)
            stem = f"batch_{bi // bsize:03d}"
            frames, finish = None, ""
            if not (saved and saved.get("failed_finish") == "RECITATION"):
                frames, finish = describe_frames(L, doc, batch, prev_titles, stem, False, CFG["max_attempts"])
            else:
                finish = "RECITATION"
            if frames is None and finish == "RECITATION":
                # Gemini's recitation filter blocks verbatim copies of published text; ask for paraphrase instead
                frames, finish = describe_frames(L, doc, batch, prev_titles, stem + "-paraphrase", True, 2)
            if frames is None:
                # last resort: one frame per request, so a single problematic frame cannot block its neighbours
                frames = []
                for c in batch:
                    one, finish = describe_frames(L, doc, [c], prev_titles, f"{stem}-{c['cid']}", True, 2)
                    frames.extend(one or [])
                missing = [c["cid"] for c in batch if c["cid"] not in {f.get("frame_id") for f in frames}]
                if missing:
                    log(f"visual {lid}: {stem} frames {missing} could not be reviewed (finish={finish})")
            write_json(bpath, {"at": now_iso(), "model": CFG["visual_model"], "candidates": batch, "frames": frames,
                               **({"failed_finish": finish} if len(frames) < len(batch) else {})})
        by_id = {f.get("frame_id"): f for f in frames}
        for c in batch:
            f = by_id.get(c["cid"])
            if f:
                results[c["cid"]] = {**c, **f}
                if f.get("useful") and not f.get("duplicate_of"):
                    prev_titles.append(f"{hms(c['t'])} - {f.get('title', '')}")
            else:
                results[c["cid"]] = {**c, "unreviewed": True}
    had_gaps = bool((st.get("visual") or {}).get("unreviewed"))
    write_visual_outputs(L, doc, cands, results)
    if had_gaps and (st.get("summary") or {}).get("done"):
        update_state(lid, lambda s: s.pop("summary", None))  # summary predates the recovered frames; rebuild it


def known_links() -> list[str]:
    urls = set()
    for L in load_catalog():
        urls.update(l["url"] for l in L["links"])
        urls.update(re.findall(r"https?://[^\s)\]>]+", L.get("description_md") or ""))
    return sorted(urls)


def verify_urls(text: str, links: list[str], also_scan: str = "") -> str:
    """Point OCR-read URLs at the exact link from a course page when they nearly match."""
    import difflib
    notes = []
    for u in re.findall(r"""https?://[^\s)\]>'"]+""", (text or "") + " " + also_scan):
        u = u.rstrip(".,;")
        if u in links:
            continue
        best = max(links, key=lambda k: difflib.SequenceMatcher(None, u, k).ratio(), default=None)
        if best and difflib.SequenceMatcher(None, u, best).ratio() >= 0.8 and best not in notes:
            notes.append(best)
    if not notes:
        return text
    return text.rstrip() + "\n\n" + "\n".join(f"*Exact link from the course page:* <{n}>" for n in notes)


def fix_urls(text: str, links: list[str]) -> str:
    """Replace URLs (with or without scheme) that nearly match a course-page link with the exact link."""
    import difflib

    def best_link(u: str) -> str | None:
        full = u if u.startswith("http") else "https://" + u
        if full in links:
            return None
        best = max(links, key=lambda k: difflib.SequenceMatcher(None, full, k).ratio(), default=None)
        if best and difflib.SequenceMatcher(None, full, best).ratio() >= 0.8:
            return best
        return None

    def repl(m: re.Match) -> str:
        raw = m.group(0)
        u = raw.rstrip(".,;")
        b = best_link(u)
        return (b + raw[len(u):]) if b else raw

    return re.sub(r"""(?:https?://|(?<![\w/.])docs\.google\.com/)[^\s)\]>'"`]+""", repl, text)


def write_visual_outputs(L: dict, doc: dict, cands: list[dict], results: dict) -> None:
    lid = L["id"]
    links = known_links()
    out = lesson_out(L)
    fdir = out / "frames"
    fdir.mkdir(parents=True, exist_ok=True)
    kept = [r for r in (results[c["cid"]] for c in cands)
            if r.get("useful") and not r.get("duplicate_of") and r.get("legibility") != "illegible"]
    unreviewed = [r for r in results.values() if r.get("unreviewed")]
    wanted = set()
    entries = []
    for r in kept:
        fname = f"{L['num']:03d}_{hms(r['t']).replace(':', '-')}.png"
        wanted.add(fname)
        dst = fdir / fname
        if not dst.exists():
            run([FFMPEG, "-y", "-v", "error", "-ss", f"{r['t']:.3f}", "-i", str(COURSE / L["video"]),
                 "-frames:v", "1", str(dst)])
        entries.append({"t": r["t"], "time": hms(r["t"]), "frame": f"frames/{fname}", "category": r.get("category"),
                        "title": r.get("title", ""), "on_screen": r.get("on_screen", ""),
                        "not_in_transcript": verify_urls(r.get("not_in_transcript") or "", links,
                                                         r.get("on_screen") or ""),
                        "legibility": r.get("legibility"),
                        "candidate_id": r["cid"]})
    for f in fdir.glob("*.png"):  # frames from an earlier run that are no longer selected
        if f.name not in wanted:
            f.unlink()
    vj = {"lesson": {k: L.get(k) for k in ("num", "id", "title", "video")}, "generated_at": now_iso(),
          "model": CFG["visual_model"], "frame_resolution": f"{L.get('width')}x{L.get('height')} (native)",
          "candidates_reviewed": len(cands) - len(unreviewed), "candidates_unreviewed": len(unreviewed),
          "frames": entries}
    write_json(out / "visual_notes.json", vj)
    write_text(out / "visual_notes.md", render_visual_md(L, vj, unreviewed))
    update_state(lid, lambda s: s.__setitem__("visual", {"done": True, "at": now_iso(), "kept": len(entries),
                                                         "reviewed": len(cands) - len(unreviewed),
                                                         "unreviewed": len(unreviewed)}))
    log(f"visual {lid}: {len(cands)} candidates -> {len(entries)} frames kept"
        + (f", {len(unreviewed)} UNREVIEWED" if unreviewed else ""))


def render_visual_md(L: dict, vj: dict, unreviewed: list[dict]) -> str:
    lines = [f"# {L['num']}. {L['title']} - Visual notes", "",
             f"- **Frames kept:** {len(vj['frames'])} of {vj['candidates_reviewed']} visually distinct candidates "
             f"(sampled from keyframes every ~5 s, de-duplicated)",
             f"- **Resolution:** {vj['frame_resolution']}, PNG in [`frames/`](frames/)",
             f"- **Described with:** {vj['model']} (Gemini API), checked against [the transcript](transcript.md)",
             "", "> **On screen** records what the frame shows. **Not in transcript** lists information that is "
             "visible but not said aloud nearby. AI-generated: verify important figures against the frame.", ""]
    if unreviewed:
        lines += [f"> Warning: {len(unreviewed)} candidate frame(s) could not be reviewed: "
                  + ", ".join(hms(u["t"]) for u in unreviewed[:20]), ""]
    if not vj["frames"]:
        lines += ["No informative on-screen content was found (presenter-only or no visual aids).", ""]
    for e in vj["frames"]:
        lines += [f"## [{e['time']}] {e['title']}", "",
                  f"![{L['num']} {e['time']} - {e['title']}]({urllib.parse.quote(e['frame'])})", "",
                  f"*{(e['category'] or '').replace('_', ' ')}*" + (f" | legibility: {e['legibility']}"
                                                                    if e.get("legibility") not in ("clear", None) else ""),
                  ""]
        if e["on_screen"]:
            lines += ["**On screen:**", "", e["on_screen"].strip(), ""]
        if e["not_in_transcript"]:
            lines += ["**Not in transcript:**", "", e["not_in_transcript"].strip(), ""]
    return "\n".join(lines)


# --------------------------------------------------------------------------- summaries

def summary_prompt(L: dict, doc: dict, vj: dict | None) -> str:
    tx = transcript_excerpt(doc, 0, 10 ** 9)
    vis = ""
    if vj and vj["frames"]:
        vis = "\n".join(f"[{e['time']}] {e['title']}: {e['on_screen']} | Not in transcript: {e['not_in_transcript']}"
                        for e in vj["frames"])
    return f"""Summarize one lesson of the "{CFG['course_name']}" course (Evolve) for a searchable study library.

Write Markdown with exactly these sections (omit a section only if it would be empty):
## Overview
2-4 sentences: what the lesson covers and who it is aimed at.
## Key teachings
Bullets. Each bullet states one claim, rule or method as the instructor teaches it, followed by a timestamp reference like [00:12:34] taken from the transcript.
## Frameworks and models
Named frameworks, step-by-step processes, classifications. Reproduce their structure faithfully.
## Numbers, benchmarks and rules of thumb
Every concrete figure the lesson gives (budgets, ROAS targets, percentages, timelines), with timestamps.
## Tools and resources mentioned
## Exercises and action items assigned
## On-screen material worth reviewing
Point to the most important frames from the visual notes by timestamp.

Rules:
- Report the course's teachings only, attributed to the instructor ("The instructor recommends..."). Do not add outside knowledge, opinions, or advice for any particular business.
- Do not invent details; if something is unclear in the transcript, say so.
- Keep it dense and scannable; no preamble.
- Plain Markdown only: no LaTeX or math notation. Write formulas in plain text, e.g. `Gross profit = revenue x margin - ad spend`.
- Copy URLs exactly as written in the lesson page text; URLs read off the screen (visual notes) may contain OCR errors.

Lesson: {L['num']}. {L['title']} (section: {L.get('section', '')}; length {hms(L.get('duration', 0))})

Lesson page text:
{L.get('description_md') or '(none)'}

Visual notes:
{vis or '(none)'}

Transcript:
{tx}
"""


def stage_summarize(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    stop = threading.Event()

    def work(L):
        if stop.is_set():
            return
        try:
            summarize_lesson(L, redo=args.redo)
        except (BudgetExceeded, QuotaExhausted) as e:
            stop.set()
            log(f"STOPPING: {e}")
        except Exception as e:
            log(f"summarize ERROR {L['id']}: {e}")

    with cf.ThreadPoolExecutor(max_workers=args.workers) as ex:
        list(ex.map(work, lessons))
    log(f"summarize finished; total recorded spend ${usage_total():.2f}")


def summarize_lesson(L: dict, redo: bool = False) -> None:
    lid = L["id"]
    st = load_state(lid)
    if st.get("summary", {}).get("done") and not redo:
        return
    doc = read_json(lesson_out(L) / "transcript.json")
    if not doc or doc["status"] != "complete":
        log(f"summarize {lid}: transcript missing or partial; skipping")
        return
    if not doc.get("alignment") and not CFG.get("allow_unaligned"):
        log(f"summarize {lid}: transcript timing not aligned yet; skipping")
        return
    vj = read_json(lesson_out(L) / "visual_notes.json")
    model = CFG["summary_model"]
    gen = {"maxOutputTokens": 16384}
    if model.startswith("gemini-3"):
        gen["thinkingConfig"] = {"thinkingLevel": CFG["thinking_level"]}
    raw_path = RAW / lid / "summary" / "summary.json"
    resp = gemini_generate(model, [{"text": summary_prompt(L, doc, vj)}], gen, raw_path,
                           {"stage": "summary", "lesson": lid, "ref": f"{lid}/summary"})
    text, finish = response_text(resp)
    if not text:
        log(f"summarize {lid}: empty response (finish={finish})")
        return
    head = [f"# {L['num']}. {L['title']} - Summary", "",
            f"> AI-generated summary of the course lesson ({model}), written from [transcript.md](transcript.md)"
            + (" and [visual notes](visual_notes.md)" if vj else "") +
            ". It describes what the course teaches; it is not a transcript and not a decision for any business. "
            "Check timestamps in the transcript before relying on a claim.", ""]
    write_text(lesson_out(L) / "summary.md", "\n".join(head) + "\n" + fix_urls(text.strip(), known_links()) + "\n")
    update_state(lid, lambda s: s.__setitem__("summary", {"done": True, "at": now_iso(), "finish": finish}))
    log(f"summarize {lid}: done ({finish})")


# --------------------------------------------------------------------------- templates

def md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [[(c or "").replace("|", "\\|").replace("\n", "<br>") for c in r] + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def extract_docx(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.search(r"heading (\d)", style)
            if m:
                out.append("#" * (int(m.group(1)) + 1) + " " + t)
            elif "title" in style:
                out.append("## " + t)
            elif "list" in style:
                out.append("- " + t)
            else:
                out.append(t)
            out.append("")
        elif tag == "tbl":
            tbl = Table(child, d)
            rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
            out += [md_table(rows), ""]
    return "\n".join(out)


def extract_xlsx(path: Path) -> str:
    import openpyxl

    wb_v = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    wb_f = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
    out = []
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        out.append(f"## Sheet: {ws.title}")
        out.append("")
        rows = []
        for rv, rf in zip(ws.iter_rows(max_col=40), wf.iter_rows(max_col=40)):
            vals = []
            for cv, cfm in zip(rv, rf):
                v = cv.value
                f = cfm.value
                if v is None and isinstance(f, str) and f.startswith("="):
                    vals.append(f"`{f}`")
                elif isinstance(f, str) and f.startswith("=") and v is not None:
                    vals.append(f"{v} (`{f}`)")
                else:
                    vals.append("" if v is None else str(v))
            if any(x.strip() for x in vals):
                rows.append(vals)
            if len(rows) >= 400:
                rows.append(["... (truncated at 400 rows)"])
                break
        if rows:
            last = max(max((i for i, x in enumerate(r) if x.strip()), default=0) for r in rows) + 1
            out.append(md_table([r[:last] for r in rows]))
        else:
            out.append("(empty sheet)")
        out.append("")
    return "\n".join(out)


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {i}")
        out.append("")
        for shp in slide.shapes:
            if shp.has_text_frame:
                t = "\n".join(p.text for p in shp.text_frame.paragraphs if p.text.strip())
                if t.strip():
                    out += [t, ""]
            if getattr(shp, "has_table", False) and shp.has_table:
                rows = [[c.text for c in r.cells] for r in shp.table.rows]
                out += [md_table(rows), ""]
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                out += [f"*Speaker notes:* {n}", ""]
    return "\n".join(out)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    r = PdfReader(str(path))
    out = []
    for i, p in enumerate(r.pages, 1):
        out += [f"## Page {i}", "", (p.extract_text() or "").strip(), ""]
    return "\n".join(out)


def template_md_path(name: str) -> Path:
    num = int(re.match(r"^(\d+)\.", name).group(1))
    stem, ext = os.path.splitext(name)
    return OUT_TEMPLATES / f"{num:03d}-{slugify(re.sub(r'^\d+\.\s*', '', stem))}{ext.lower()}.md"


def stage_templates(args) -> None:
    lessons = select_lessons(args.lessons)
    by_name = {}
    for L in lessons:
        for t in L["templates"]:
            by_name[t] = L
    for name, L in by_name.items():
        src = COURSE / name
        dst = template_md_path(name)
        ext = src.suffix.lower()
        try:
            body = {".docx": extract_docx, ".xlsx": extract_xlsx, ".pptx": extract_pptx, ".pdf": extract_pdf}[ext](src)
        except Exception as e:
            body = f"(text extraction failed: {e})"
        head = [f"# Template text: {name}", "",
                f"- **Original file:** [{name}]({rel_link(src, OUT_TEMPLATES)}) (open this for formatting, formulas "
                "and images)",
                f"- **Lesson:** [{L['num']}. {L['title']}]({rel_link(lesson_out(L) / 'lesson_page.md', OUT_TEMPLATES)})",
                "", "> Plain-text extract for searching. The original file in the course folder is unchanged.", ""]
        write_text(dst, "\n".join(head) + "\n" + body.strip() + "\n")
        log(f"template {name}: {len(body)} chars")


# --------------------------------------------------------------------------- lesson pages, index, report

def write_lesson_page(L: dict) -> None:
    out = lesson_out(L)
    lines = [f"# {L['num']}. {L['title']}", "", f"- **Course section:** {L.get('section', '')}"]
    if L.get("skool_url"):
        lines.append(f"- **Skool lesson:** <{L['skool_url']}>")
    if L.get("html"):
        lines.append(f"- **Saved course page:** [{L['html']}]({rel_link(COURSE / L['html'], out)})")
    if L.get("video"):
        lines.append(f"- **Video:** [{L['video']}]({rel_link(COURSE / L['video'], out)}) ({hms(L['duration'])})")
    for t in L["templates"]:
        lines.append(f"- **Template:** [{t}]({rel_link(COURSE / t, out)}) - "
                     f"[text extract]({rel_link(template_md_path(t), out)})")
    lines += ["", "## Lesson page text", "", L.get("description_md") or "*(The course page has no text "
              "besides the video.)*", ""]
    write_text(out / "lesson_page.md", "\n".join(lines))


def stage_index(args) -> None:
    cat = read_json(CATALOG)
    lessons = cat["lessons"]
    for L in lessons:
        write_lesson_page(L)
    lines = [f"# {CFG['course_name']} - Course index", "",
             f"Generated {now_iso()} from `{COURSE.name}`. See [README.md](README.md) for how this library was "
             "built and [PROCESSING_REPORT.md](PROCESSING_REPORT.md) for completeness and flagged sections.", "",
             "Legend: **T** transcript (Markdown / JSON / SRT), **V** visual notes + screenshots, **S** summary, "
             "**P** lesson page text.", ""]
    sections: dict[str, list[dict]] = {}
    for L in sorted(lessons, key=lambda x: (x.get("section_order", 99), x["num"])):
        sections.setdefault(L.get("section") or "Unsorted", []).append(L)
    lines.append("## Contents")
    lines.append("")
    for s in sections:
        lines.append(f"- [{s}](#{slugify(s, 200)}) ({len(sections[s])} lessons)")
    lines.append("")
    for s, ls in sections.items():
        lines += [f'<a id="{slugify(s, 200)}"></a>', "", f"## {s}", ""]
        for L in ls:
            out = lesson_out(L)
            rl = lambda p: rel_link(p, LIB)
            st = load_state(L["id"])
            dur = f" - {hms(L['duration'])}" if L.get("video") else " - no video"
            lines.append(f"### {L['num']}. {L['title']}{dur}")
            lines.append("")
            if L.get("video"):
                tj = read_json(out / "transcript.json")
                if tj:
                    flagged = len([i for i in tj["issues"] if i["type"] not in ("transcriber_note", "unclear_words")])
                    status = "complete" if tj["status"] == "complete" else "**partial**"
                    flag = f", {flagged} flagged section(s)" if flagged else ""
                    lines.append(f"- **T:** [Markdown]({rl(out / 'transcript.md')}) | [JSON]({rl(out / 'transcript.json')}) "
                                 f"| [SRT]({rl(out / 'transcript.srt')}) - {status}{flag}")
                else:
                    lines.append("- **T:** not transcribed yet")
                vj = read_json(out / "visual_notes.json")
                if vj:
                    lines.append(f"- **V:** [Visual notes]({rl(out / 'visual_notes.md')}) - {len(vj['frames'])} "
                                 f"screenshot(s) in [frames/]({rl(out / 'frames')})")
                else:
                    lines.append("- **V:** not processed yet")
                if (out / "summary.md").exists():
                    lines.append(f"- **S:** [Summary]({rl(out / 'summary.md')})")
            lines.append(f"- **P:** [Lesson page]({rl(out / 'lesson_page.md')})"
                         + (f" | [Skool]({L['skool_url']})" if L.get("skool_url") else ""))
            for t in L["templates"]:
                lines.append(f"- **Template:** [{t}]({rl(COURSE / t)}) ([text]({rl(template_md_path(t))}))")
            ext = [l for l in L["links"] if "skool.com" not in l["url"]]
            if ext:
                lines.append("- **Linked resources:** " + " | ".join(
                    f"[{(l['text'] or l['url'])[:60]}]({l['url']})" for l in ext))
            lines.append("")
    write_text(LIB / "INDEX.md", "\n".join(lines))
    write_report(lessons)
    write_search(lessons)
    log("index: INDEX.md, PROCESSING_REPORT.md, search.html written")


def write_report(lessons: list[dict]) -> None:
    vids = [L for L in lessons if L.get("video")]
    rows, detail = [], []
    tot = {"chunks": 0, "done": 0, "uncertain": 0, "failed": 0, "pending": 0}
    for L in vids:
        st = load_state(L["id"])
        prep = st.get("prepare") or {}
        n = len(prep.get("chunks", []))
        tr = st.get("transcribe", {})
        c = {k: sum(1 for v in tr.values() if v.get("status") == k) for k in ("done", "uncertain", "failed")}
        pend = n - sum(c.values())
        tot["chunks"] += n
        tot["pending"] += pend
        for k in c:
            tot[k] += c[k]
        tj = read_json(lesson_out(L) / "transcript.json")
        vj = read_json(lesson_out(L) / "visual_notes.json")
        iss = tj["issues"] if tj else []
        cnt = lambda t: sum(1 for i in iss if i["type"] == t)
        al = (tj or {}).get("alignment") or {}
        timing = f"{al['match_ratio']:.0%}" if al else ("-" if not tj else "**no**")
        rows.append(f"| {L['num']} | {L['title'][:48]} | {hms(L['duration'])} | {c['done']}/{n} | {c['uncertain']} | "
                    f"{c['failed']} | {pend} | {timing} | {cnt('possible_missing_speech') + cnt('possible_omission')} | "
                    f"{al.get('whisper_fill_words', 0) if al else '-'} | {al.get('removed_words', 0) if al else '-'} | "
                    f"{cnt('unverified_passage')} | {cnt('unclear_words')} | "
                    f"{(str(len(vj['frames'])) if vj else '-')} | {'yes' if (lesson_out(L) / 'summary.md').exists() else '-'} |")
        serious = [i for i in iss if i["type"] in ("missing_chunk", "not_transcribed_yet", "uncertain_chunk",
                                                    "possible_missing_speech", "low_word_rate", "hard_chunk_boundary",
                                                    "possible_omission", "unverified_passage", "pending_audio_check",
                                                    "filled_omission", "removed_unspoken_text")]
        if serious or (vj and vj.get("candidates_unreviewed")):
            detail += [f"### {L['num']}. {L['title']}", ""]
            for i in serious:
                rng = hms(i["start"]) if i["start"] == i["end"] else f"{hms(i['start'])}-{hms(i['end'])}"
                detail.append(f"- `[{rng}]` {ISSUE_LABEL.get(i['type'], i['type'])}: {i['detail']}")
            if vj and vj.get("candidates_unreviewed"):
                detail.append(f"- {vj['candidates_unreviewed']} candidate frame(s) not reviewed (visual batch failed)")
            detail.append("")
    spend = usage_total()
    lines = ["# Processing report", "", f"Updated {now_iso()}.", "",
             f"- Videos: {len(vids)} | audio chunks: {tot['chunks']} | done: {tot['done']} | uncertain: "
             f"{tot['uncertain']} | failed: {tot['failed']} | pending: {tot['pending']}",
             f"- Recorded API spend: ${spend:.2f} (from `_pipeline/state/usage.jsonl`, standard paid-tier prices)",
             "", "**Column key:** *Chunks* = transcribed / total; *Uncertain* = best of several attempts still failed a "
             "quality check (kept, flagged); *Failed* = no usable transcription (section missing); *Timing* = share "
             "of Gemini words matched to local Whisper word timestamps (\"no\" = not aligned yet); *Omissions* = "
             "speech Whisper heard that the transcript lacks (low confidence, not filled); *Filled* = words inserted "
             "from Whisper, marked *[+ ...]*; *Removed* = words Gemini wrote that two local checks found are not in "
             "the audio (kept in transcript.json); *Unconfirmed* = passages neither local check could confirm "
             "(usually crosstalk, music, played ads or foreign speech; kept, flagged); *Unclear* = segments with "
             "[inaudible] or [?].", "",
             "| # | Lesson | Length | Chunks | Uncertain | Failed | Pending | Timing | Omissions | Filled | Removed "
             "| Unconfirmed | Unclear | Frames | Summary |",
             "|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|"] + rows + ["", "## Sections needing attention", ""]
    lines += detail or ["None."]
    write_text(LIB / "PROCESSING_REPORT.md", "\n".join(lines) + "\n")


def write_search(lessons: list[dict]) -> None:
    items = []
    for L in lessons:
        out = lesson_out(L)
        base = rel_link(out, LIB)
        tj = read_json(out / "transcript.json")
        if tj:
            for s in tj["segments"]:
                items.append([L["num"], round(s["start"]), "t", s["text"]])
        vj = read_json(out / "visual_notes.json")
        if vj:
            for e in vj["frames"]:
                items.append([L["num"], round(e["t"]), "v",
                              f"{e['title']}. {e['on_screen']} {e['not_in_transcript']}".strip()])
        if L.get("description_md"):
            items.append([L["num"], -1, "p", L["description_md"]])
    meta = {L["num"]: {"title": L["title"], "dir": rel_link(lesson_out(L), LIB), "section": L.get("section", "")}
            for L in lessons}
    data = json.dumps({"meta": meta, "items": items}, ensure_ascii=False).replace("</", "<\\/")
    page = SEARCH_HTML.replace("__DATA__", data).replace("__TITLE__", html.escape(CFG["course_name"]))
    write_text(LIB / "search.html", page)


SEARCH_HTML = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Course Search</title>
<style>
:root{--bg:#fbfbf9;--fg:#1d1d1b;--muted:#6b6b66;--line:#e4e3de;--accent:#2457c5;--mark:#ffe58a}
@media (prefers-color-scheme:dark){:root{--bg:#161615;--fg:#ecebe6;--muted:#9a9993;--line:#2c2c2a;--accent:#8cb0ff;--mark:#6b5a12}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.5 system-ui,sans-serif}
main{max-width:900px;margin:0 auto;padding:24px 16px}h1{font-size:20px;margin:0 0 4px}p.sub{color:var(--muted);margin:0 0 16px}
.bar{display:flex;gap:8px;flex-wrap:wrap}input[type=search]{flex:1 1 260px;font:inherit;padding:10px 12px;border:1px solid var(--line);border-radius:8px;background:transparent;color:inherit}
label{color:var(--muted);font-size:14px;display:flex;gap:4px;align-items:center}
.r{border-bottom:1px solid var(--line);padding:10px 0}.r a{color:var(--accent);text-decoration:none}.k{color:var(--muted);font-size:13px}
mark{background:var(--mark);color:inherit}#n{color:var(--muted);font-size:14px;margin:12px 0}
</style></head><body><main>
<h1>__TITLE__ - search</h1><p class="sub">Searches transcripts, on-screen notes and lesson pages. All words must match.</p>
<div class="bar"><input id="q" type="search" placeholder="e.g. ROAS target, bundle offer, Andromeda" autofocus>
<label><input type="checkbox" id="ft" checked>transcripts</label><label><input type="checkbox" id="fv" checked>on-screen</label><label><input type="checkbox" id="fp" checked>pages</label></div>
<div id="n"></div><div id="out"></div></main>
<script>
const D=__DATA__;const q=document.getElementById('q'),out=document.getElementById('out'),n=document.getElementById('n');
const pad=x=>String(x).padStart(2,'0');const hms=t=>t<0?'':`${pad(Math.floor(t/3600))}:${pad(Math.floor(t%3600/60))}:${pad(t%60)}`;
const esc=s=>s.replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const low=D.items.map(i=>i[3].toLowerCase());
function run(){const terms=q.value.toLowerCase().split(/\\s+/).filter(Boolean);out.innerHTML='';n.textContent='';if(!terms.length)return;
const on={t:ft.checked,v:fv.checked,p:fp.checked};const res=[];
for(let i=0;i<D.items.length&&res.length<400;i++){const it=D.items[i];if(!on[it[2]])continue;if(terms.every(t=>low[i].includes(t)))res.push(it)}
n.textContent=res.length>=400?'First 400 matches':res.length+' match(es)';
const re=new RegExp('('+terms.map(t=>t.replace(/[.*+?^${}()|[\\]\\\\]/g,'\\\\$&')).join('|')+')','gi');
out.innerHTML=res.map(([num,t,k,txt])=>{const m=D.meta[num];const f=k==='t'?'transcript.md':k==='v'?'visual_notes.md':'lesson_page.md';
const lab=k==='t'?'transcript':k==='v'?'on-screen':'page';let s=txt;if(s.length>420){const p=Math.max(0,s.toLowerCase().indexOf(terms[0])-160);s=(p?'... ':'')+s.slice(p,p+420)+' ...'}
return `<div class="r"><div class="k">${esc(m.section)} &middot; ${lab} ${hms(t)}</div><a href="${m.dir}/${f}">${num}. ${esc(m.title)}</a><div>${esc(s).replace(re,'<mark>$1</mark>')}</div></div>`}).join('')}
let tm;q.addEventListener('input',()=>{clearTimeout(tm);tm=setTimeout(run,150)});[ft,fv,fp].forEach(e=>e.addEventListener('change',run));
</script></body></html>
"""


# --------------------------------------------------------------------------- follow (chains the later stages)

CHECK_LOCK = threading.Lock()


def unchecked(L: dict) -> list:
    return read_json(lesson_dir_state(L["id"]) / "audio_checks_pending.json", []) or []


def run_audio_checks(L: dict) -> None:
    """small.en re-check of suspected inserted text (runs whisper_timing.py --verify in its own Python)."""
    uv = shutil.which("uv") or "uv"
    with CHECK_LOCK:  # one small.en model at a time on this CPU
        p = subprocess.run([uv, "run", "--python", "3.14", "--script", str(HERE / "whisper_timing.py"), "--verify",
                            "--lessons", str(L["num"])], capture_output=True, text=True, encoding="utf-8",
                           errors="replace", env={**os.environ, "HF_HUB_DISABLE_SYMLINKS_WARNING": "1"})
    for line in p.stdout.splitlines():
        if line.startswith("check "):
            log(f"  {line}")
    if p.returncode != 0:
        log(f"audio check FAILED {L['id']}: {p.stderr[-800:]}")


def stage_checks(args) -> None:
    for L in select_lessons(args.lessons, video_only=True):
        if unchecked(L):
            run_audio_checks(L)
            assemble_lesson(L)


def stage_follow(args) -> None:
    """Assemble, describe and summarize each lesson as soon as its transcript and Whisper timing are ready.

    Polls every 3 minutes until every video lesson is finished (or nothing more can progress)."""
    stop = threading.Event()

    def process(L):
        if stop.is_set():
            return
        try:
            asm = load_state(L["id"]).get("assemble") or {}
            if not asm.get("aligned") or asm.get("status") != "complete" or unchecked(L):
                assemble_lesson(L)
            if unchecked(L):
                run_audio_checks(L)
                assemble_lesson(L)
            if unchecked(L):
                log(f"follow {L['id']}: audio re-check incomplete; visual/summary postponed")
                return
            visual_lesson(L)
            summarize_lesson(L)
        except (BudgetExceeded, QuotaExhausted) as e:
            stop.set()
            log(f"STOPPING: {e}")
        except Exception as e:
            log(f"follow ERROR {L['id']}: {e}\n{traceback.format_exc(limit=3)}")

    while not stop.is_set():
        ready, waiting, todo = [], [], []
        for L in select_lessons(args.lessons, video_only=True):
            st = load_state(L["id"])
            n = len((st.get("prepare") or {}).get("chunks", []))
            tr = st.get("transcribe", {})
            tx_done = n > 0 and all(tr.get(str(i), {}).get("status") in ("done", "uncertain", "failed")
                                    for i in range(n))
            timed = (lesson_dir_state(L["id"]) / CFG.get("whisper_file", "whisper_words.json")).exists()
            if not (tx_done and timed):
                waiting.append(L["id"])
                continue
            vis = st.get("visual") or {}
            finished = (st.get("assemble") or {}).get("aligned") and vis.get("done") and not vis.get("unreviewed") \
                and (st.get("summary") or {}).get("done")
            (ready if finished else todo).append(L)
        if todo:
            log(f"follow: processing {len(todo)} lesson(s); {len(waiting)} waiting for transcription/Whisper")
            with cf.ThreadPoolExecutor(max_workers=args.workers) as ex:
                list(ex.map(process, todo))
            stage_index(args)
        if not waiting or stop.is_set():
            break
        log(f"follow: {len(ready) + len(todo)} ready, {len(waiting)} waiting; next check in 3 min")
        time.sleep(180)
    stage_index(args)
    log(f"follow finished; total recorded spend ${usage_total():.2f}")


# --------------------------------------------------------------------------- estimate / status

def stage_estimate(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    audio_s = speech_s = 0.0
    n_chunks = n_cands = 0
    missing_prep = missing_frames = 0
    for L in lessons:
        st = load_state(L["id"])
        prep = st.get("prepare")
        if prep:
            audio_s += prep["audio_duration"]
            speech_s += prep["speech_seconds"]
            n_chunks += len(prep["chunks"])
        else:
            missing_prep += 1
            audio_s += L["duration"]
            speech_s += L["duration"] * 0.85
            n_chunks += math.ceil(L["duration"] / CFG["chunk_target_s"])
        fr = st.get("frames")
        if fr:
            n_cands += len(fr["candidates"])
        else:
            missing_frames += 1
            n_cands += int(L["duration"] / 25)
    pr = CFG["prices"][CFG["transcribe_model"]]
    words = speech_s / 60 * 165
    # transcription: audio tokens + prompt; output = text + JSON/timestamp overhead + low thinking
    t_in_audio = audio_s * 32
    t_in_text = n_chunks * 900
    t_out = words * 1.35 + (words / 14) * 22
    t_think = 0.25 * t_out
    c_tr = (t_in_audio * pr["audio"] + t_in_text * pr["input"] + (t_out + t_think) * pr["output"]) / 1e6
    # visual: images (assume <= 1100 tokens each) + transcript excerpts + prompt; ~200 output tokens/frame
    batches = math.ceil(n_cands / CFG["visual_batch"]) if n_cands else 0
    v_in = n_cands * 1100 + batches * (1800 + 1200)
    v_out = n_cands * 200 * 1.3
    c_vis = (v_in * pr["input"] + v_out * pr["output"]) / 1e6
    # summaries: full transcript once + visual notes; ~1800 output tokens
    s_in = words * 1.35 * 1.25 + len(lessons) * 1500
    s_out = len(lessons) * 1800 * 1.3
    c_sum = (s_in * pr["input"] + s_out * pr["output"]) / 1e6
    retry = 0.20
    print(f"Lessons with video: {len(lessons)}  (not yet prepared: {missing_prep}, frames not sampled: {missing_frames})")
    print(f"Audio: {audio_s / 3600:.2f} h  speech: {speech_s / 3600:.2f} h  chunks: {n_chunks}  frame candidates: {n_cands}")
    print(f"Model: {CFG['transcribe_model']}  prices/1M tokens: input ${pr['input']}, audio ${pr['audio']}, output ${pr['output']}")
    print(f"  transcription : in {t_in_audio / 1e6:.2f}M audio + {t_in_text / 1e6:.2f}M text tok, out ~{(t_out + t_think) / 1e6:.2f}M tok  -> ${c_tr:.2f}")
    print(f"  visual notes  : in ~{v_in / 1e6:.2f}M tok ({n_cands} images), out ~{v_out / 1e6:.2f}M tok          -> ${c_vis:.2f}")
    print(f"  summaries     : in ~{s_in / 1e6:.2f}M tok, out ~{s_out / 1e6:.2f}M tok                        -> ${c_sum:.2f}")
    total = c_tr + c_vis + c_sum
    print(f"  TOTAL ~${total:.2f}  (with {int(retry * 100)}% retry margin: ~${total * (1 + retry):.2f}); "
          f"spend already recorded: ${usage_total():.2f}; hard cap max_cost_usd=${CFG['max_cost_usd']}")


def stage_status(args) -> None:
    lessons = select_lessons(args.lessons, video_only=True)
    print(f"{'#':>3}  {'chunks':>9}  {'unc':>3} {'fail':>4}  {'asm':>8}  {'cand':>4} {'kept':>4}  sum  title")
    for L in lessons:
        st = load_state(L["id"])
        n = len((st.get("prepare") or {}).get("chunks", []))
        tr = st.get("transcribe", {})
        d = sum(1 for v in tr.values() if v.get("status") in ("done", "uncertain"))
        u = sum(1 for v in tr.values() if v.get("status") == "uncertain")
        f = sum(1 for v in tr.values() if v.get("status") == "failed")
        a = (st.get("assemble") or {}).get("status", "-")
        c = len((st.get("frames") or {}).get("candidates", [])) if st.get("frames") else "-"
        k = (st.get("visual") or {}).get("kept", "-")
        s = "yes" if (st.get("summary") or {}).get("done") else "-"
        print(f"{L['num']:>3}  {d:>4}/{n:<4}  {u:>3} {f:>4}  {a:>8}  {c!s:>4} {k!s:>4}  {s:>3}  {L['title'][:50]}")
    print(f"recorded spend: ${usage_total():.2f}")


# --------------------------------------------------------------------------- main

STAGES = {
    "inventory": stage_inventory, "prepare": stage_prepare, "transcribe": stage_transcribe,
    "assemble": stage_assemble, "frames": stage_frames, "visual": stage_visual, "summarize": stage_summarize,
    "templates": stage_templates, "index": stage_index, "estimate": stage_estimate, "status": stage_status,
    "follow": stage_follow, "checks": stage_checks,
}
ALL_ORDER = ["inventory", "prepare", "transcribe", "assemble", "frames", "visual", "summarize", "templates", "index"]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("stage", choices=list(STAGES) + ["all", "free"])
    ap.add_argument("--lessons", help="lesson numbers, e.g. 20 or 1,5,10-12 (default: all)")
    ap.add_argument("--workers", type=int, default=CFG["workers"])
    ap.add_argument("--redo", action="store_true", help="redo work even if already complete")
    ap.add_argument("--retry-failed", action="store_true", help="retry chunks that previously failed")
    ap.add_argument("--whisper-file", help="timing file name in each lesson's state dir (default whisper_words.json)")
    args = ap.parse_args()
    if args.whisper_file:
        CFG["whisper_file"] = args.whisper_file
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if args.stage == "all":
        order = ALL_ORDER
    elif args.stage == "free":
        order = ["inventory", "prepare", "frames", "templates", "index"]
    else:
        order = [args.stage]
    for s in order:
        if s == "inventory" and CATALOG.exists() and args.stage in ("all", "free") and not args.redo:
            continue
        STAGES[s](args)


if __name__ == "__main__":
    main()
